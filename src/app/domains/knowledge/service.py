import asyncio
import fnmatch
import io
import json
import math
import re
import traceback
import xml.etree.ElementTree as ET
from collections import deque
from collections.abc import Awaitable, Callable
from dataclasses import dataclass
from typing import Any
from urllib.parse import urljoin, urlparse, urlunparse
from uuid import UUID

import httpx
import structlog
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession


def _parse_html(html: str) -> BeautifulSoup:
    """
    Prefer lxml over ``html.parser``: many production sites (WordPress, legacy CMS)
    ship broken conditional comments or tags that confuse the stdlib parser into an
    empty or truncated DOM, so link discovery and ``get_text`` silently see nothing.
    """
    return BeautifulSoup(html, "lxml")

from app.core.errors import AppError
from app.core.settings import get_settings
from app.domains.plans.plan_limits import (
    DEFAULT_KNOWLEDGE_STORAGE_CAP_BYTES,
    STARTER_KNOWLEDGE_STORAGE_CAP_BYTES,
    training_storage_cap_bytes,
    website_crawl_cap_bytes,
)
from app.domains.knowledge.schemas import (
    FileSourceListItemDTO,
    FileUploadResultDTO,
    IndexJobDTO,
    KnowledgeSourceCreateRequest,
    KnowledgeSourceDTO,
    QAPairDetailDTO,
    QAPairListItemDTO,
    TextSnippetDetailDTO,
    TextSnippetListItemDTO,
    WebsiteIndividualRequest,
    WebsiteIngestBase,
    WebsiteMode,
    WebsitePathRule,
    WebsiteSourceListItemDTO,
    WebsiteUrlPreviewRequest,
    WebsiteUrlPreviewResponse,
    WebsiteUsageResponse,
)


@dataclass
class SitemapDiscoveryResult:
    """Diagnostics from ``_discover_sitemap_urls`` (empty ``urls`` can mean filters, fetch failure, or bad XML)."""

    urls: list[str]
    http_success_count: int = 0
    http_failure_count: int = 0
    xml_parse_failure_count: int = 0
    had_successful_xml_document: bool = False
    truncated_by_doc_cap: bool = False
    urls_discovered_total: int = 0
    urls_after_filter: int = 0
    hub_urls_added: int = 0


EMBEDDING_DIMENSION = 1536
# Hard cap on how many distinct page URLs one crawl job may visit (safety rail; byte budget is primary).
MAX_CRAWL_PAGES_SAFETY_CEILING = 2000
MAX_DASHBOARD_WEBSITE_PAGES = MAX_CRAWL_PAGES_SAFETY_CEILING
# With include path rules, follow same-site links that pass excludes only so hub/category pages can lead to matching URLs.
DASHBOARD_BFS_INCLUDE_RULE_MAX_HOPS = 14
DASHBOARD_BFS_INCLUDE_RULE_MAX_PAGES_VISITED = 3500
# Per-response byte charge cap for crawl HTTP bodies (in addition to total indexed storage cap).
_CRAWL_BODY_CHARGE_CAP_BYTES = 400_000
STARTER_HIDDEN_STORAGE_GRACE_RATIO = 0.23
# OpenAI embeddings cap is ~300k tokens per request; batch conservatively (tiktoken can exceed char/4).
_EMBED_BATCH_MAX_TOKENS_EST = 200_000
_EMBED_BATCH_MAX_INPUTS = 2048
_CHUNK_PERSIST_BATCH_SIZE = 32

log = structlog.get_logger("knowledge.service")


def _charge_bytes_and_html_from_response(response: httpx.Response, *, remaining_budget_bytes: int) -> tuple[int, str]:
    """Charge and parse only bytes that still fit within remaining crawl budget."""
    raw = response.content or b""
    n = len(raw)
    if remaining_budget_bytes <= 0:
        return 0, ""
    cap = min(_CRAWL_BODY_CHARGE_CAP_BYTES, remaining_budget_bytes)
    charged = min(n, cap)
    chunk = raw if n <= cap else raw[:cap]
    enc = response.encoding or "utf-8"
    try:
        html = chunk.decode(enc, errors="replace")
    except (LookupError, UnicodeDecodeError):
        html = chunk.decode("utf-8", errors="replace")
    return charged, html


# Dashboard crawl: sitemap seeds URL rows in bulk; HTML fetch runs in the worker (not HTTP rate limits).
# Page rows are upserted in chunks after fetch — not one DB round-trip per page during crawl.
DASHBOARD_CRAWL_CONTENT_BATCH = 24
DASHBOARD_CRAWL_PROGRESS_EVERY = 12
DASHBOARD_PAGE_UPSERT_CHUNK = 100

# Large shops can expose thousands of nested sitemap index URLs; cap GETs so discovery finishes.
SITEMAP_MAX_DOCUMENT_FETCHES = 500
# Concurrent GETs when walking nested sitemap indexes (same origin session).
SITEMAP_PARALLEL_FETCHES = 8
# Collect every ``<url><loc>`` from sitemaps before path filters (Shopify catalogs are often 2k+).
SITEMAP_MAX_PAGE_URLS_COLLECTED = 50_000
# After path filters match hub pages (collections, categories), fetch HTML for embedded product paths.
HUB_EXPAND_MAX_HUB_VISITS = 64
HUB_EXPAND_MAX_LINKS_PER_HUB = 1200
# Batch multi-row inserts when seeding sitemap URL placeholders for the dashboard.
DASHBOARD_SEED_URL_INSERT_CHUNK = 250

# Many sites return a minimal shell or challenge to non-browser clients; match typical browser fetch.
_WEBSITE_CRAWL_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/122.0.0.0 Safari/537.36 ChatRelyIndexer/1.0"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
}


def _normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _safe_storage_path(filename: str) -> str:
    clean = re.sub(r"[^a-zA-Z0-9._-]+", "_", (filename or "upload").strip())[:140]
    if not clean:
        clean = "upload"
    return f"inline/{clean}"


def _extract_text_from_file_bytes(filename: str, content_type: str | None, payload: bytes) -> str:
    name = (filename or "").strip()
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    ctype = (content_type or "").lower()

    if ext == "txt" or ctype.startswith("text/plain"):
        return payload.decode("utf-8", errors="replace")

    if ext == "pdf" or ctype == "application/pdf":
        reader = PdfReader(io.BytesIO(payload))
        parts: list[str] = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                continue
        return "\n\n".join(parts)

    if ext == "docx" or "wordprocessingml.document" in ctype:
        doc = Document(io.BytesIO(payload))
        return "\n".join(p.text for p in doc.paragraphs if p.text)

    if ext == "doc" or ctype == "application/msword":
        # Legacy .doc is often either RTF content or binary OLE. Use best-effort
        # text extraction so indexing can proceed for common documents.
        if payload.lstrip().startswith(b"{\\rtf"):
            rtf = payload.decode("latin-1", errors="ignore")
            text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", rtf)
            text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
            text = text.replace("{", " ").replace("}", " ")
            return text

        # Binary fallback: decode and keep readable runs.
        candidates = [
            payload.decode("utf-16le", errors="ignore"),
            payload.decode("utf-8", errors="ignore"),
            payload.decode("latin-1", errors="ignore"),
        ]
        for candidate in candidates:
            parts = re.findall(r"[A-Za-z0-9][A-Za-z0-9\s,.;:!?()/'\"@#%&*\-]{4,}", candidate)
            merged = "\n".join(p.strip() for p in parts if p.strip())
            if merged.strip():
                return merged

    raise AppError(
        code="knowledge.file_type_unsupported",
        message="Unsupported file type. Use .pdf, .txt, .doc, or .docx",
        status_code=422,
    )


def _truncate_utf8_to_bytes(text: str, max_bytes: int) -> str:
    """Return text truncated to max UTF-8 bytes without splitting multibyte chars."""
    if max_bytes <= 0 or not text:
        return ""
    raw = text.encode("utf-8")
    if len(raw) <= max_bytes:
        return text
    return raw[:max_bytes].decode("utf-8", errors="ignore")


def _local_xml_tag(tag: str) -> str:
    return tag.split("}", 1)[-1] if "}" in tag else tag


def _is_markdown_heading(line: str) -> bool:
    return bool(re.match(r"^\s{0,3}#{1,6}\s+\S", line or ""))


def _extract_heading_path_from_line(line: str) -> tuple[int, str] | None:
    m = re.match(r"^\s{0,3}(#{1,6})\s+(.+)$", line or "")
    if not m:
        return None
    level = len(m.group(1))
    title = _normalize_text(m.group(2))
    if not title:
        return None
    return level, title


def _looks_like_table_block(block: str) -> bool:
    lines = [ln.strip() for ln in block.splitlines() if ln.strip()]
    if len(lines) < 2:
        return False
    if "|" not in lines[0] or "|" not in lines[1]:
        return False
    divider = lines[1].replace("|", "").replace(":", "").replace("-", "").strip()
    return divider == ""


def _split_large_plaintext_block(block: str, chunk_size: int, overlap: int) -> list[str]:
    pieces: list[str] = []
    start = 0
    while start < len(block):
        end = min(start + chunk_size, len(block))
        piece = block[start:end].strip()
        if piece:
            pieces.append(piece)
        if end >= len(block):
            break
        start = max(0, end - overlap)
    return pieces


def _chunk_text(text_value: str, chunk_size: int = 3600, overlap: int = 500) -> list[str]:
    """
    Semantic-first chunking:
    - split by paragraph blocks
    - carry active heading path into every chunk
    - keep Markdown code fences and tables atomic when possible
    """
    if not text_value:
        return []
    blocks = [b.strip() for b in re.split(r"\n{2,}", text_value) if b.strip()]
    if not blocks:
        return []

    chunks: list[str] = []
    current = ""
    heading_stack: dict[int, str] = {}

    def heading_prefix() -> str:
        if not heading_stack:
            return ""
        ordered = [heading_stack[level] for level in sorted(heading_stack.keys())]
        return " > ".join([h for h in ordered if h])

    def flush_current() -> None:
        nonlocal current
        if current.strip():
            chunks.append(current.strip())
        current = ""

    for raw_block in blocks:
        block = raw_block.strip()
        if not block:
            continue

        first_line = block.splitlines()[0].strip()
        heading_match = _extract_heading_path_from_line(first_line)
        if heading_match:
            level, title = heading_match
            heading_stack = {k: v for k, v in heading_stack.items() if k < level}
            heading_stack[level] = title

        prefix = heading_prefix()
        decorated_block = block
        if prefix and not _is_markdown_heading(first_line):
            decorated_block = f"Section: {prefix}\n\n{block}"

        candidate = f"{current}\n\n{decorated_block}".strip() if current else decorated_block
        if len(candidate) <= chunk_size:
            current = candidate
            continue

        flush_current()

        block_is_atomic = (
            decorated_block.startswith("```")
            or "```" in decorated_block
            or _looks_like_table_block(decorated_block)
        )
        if block_is_atomic or len(decorated_block) <= chunk_size:
            current = decorated_block
            continue

        for piece in _split_large_plaintext_block(decorated_block, chunk_size=chunk_size, overlap=overlap):
            chunks.append(piece)

    flush_current()
    return chunks


def _token_estimate(chunk: str) -> int:
    return max(1, math.ceil(len(chunk.split()) * 1.3))


def _approx_embed_request_tokens(chunk: str) -> int:
    """Conservative per-input token estimate for OpenAI embedding batch sizing."""
    if not chunk:
        return 1
    # Dense markup, URLs, or non‑Latin text can be far below 4 chars/token; /2 is a safe upper-ish bound.
    char_est = math.ceil(len(chunk) / 2)
    word_est = max(1, math.ceil(len(chunk.split()) * 1.35))
    return max(char_est, word_est)


def _normalize_url_string(url: str) -> str:
    """Canonical page URL for indexing and path-rule matching (no query or fragment)."""
    parsed = urlparse(url.strip())
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise AppError(code="validation.invalid_input", message="Invalid URL", status_code=422)
    return f"{parsed.scheme}://{parsed.netloc}{parsed.path or '/'}"


def _normalize_sitemap_document_url(url: str) -> str:
    """Fetch URL for nested sitemap XML; preserves query (Shopify ``?from=&to=`` child indexes)."""
    parsed = urlparse(url.strip())
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise AppError(code="validation.invalid_input", message="Invalid URL", status_code=422)
    return urlunparse(
        (
            parsed.scheme,
            parsed.netloc,
            parsed.path or "/",
            "",
            parsed.query,
            "",
        )
    )


def normalize_dashboard_website_url(protocol: str, url_input: str) -> str:
    raw = url_input.strip()
    if raw.lower().startswith("http://") or raw.lower().startswith("https://"):
        return _normalize_url_string(raw)
    path = raw.lstrip("/")
    combined = f"{protocol}{path}"
    return _normalize_url_string(combined)


def _url_path_for_rules(url: str) -> str:
    """Path rules use ``urlparse(url).path`` only (leading slash, no query or fragment), casefolded when matching."""
    parsed = urlparse(url)
    return parsed.path or "/"


def _rule_matches(operator: str, pattern: str, path_value: str) -> bool:
    """Path rules are matched case-insensitively (URLs often differ only by ``And`` vs ``and``)."""
    pv = path_value.casefold()
    pat = pattern.casefold()
    if operator == "starts_with":
        return pv.startswith(pat)
    if operator == "ends_with":
        return pv.endswith(pat)
    if operator == "contains":
        return pat in pv
    if operator == "exact_match":
        return pv == pat
    if operator == "wildcard":
        return fnmatch.fnmatch(pv, pat)
    return False


_NON_HTML_PAGE_PATH_SUFFIXES = (".atom", ".oembed", ".json", ".xml", ".js", ".css", ".map")


def _is_indexable_html_page_url(url: str) -> bool:
    """Drop feed/alternate format URLs often linked from Shopify collection pages."""
    path = _url_path_for_rules(url).lower()
    return not any(path.endswith(suffix) for suffix in _NON_HTML_PAGE_PATH_SUFFIXES)


def _url_passes_filters(url: str, include_rules: list[dict[str, str]], exclude_rules: list[dict[str, str]]) -> bool:
    """Exclude rules run first (any match drops the URL). With no include rules, URL passes if not excluded.

    Multiple include chips are OR: the path must match **at least one** include rule when includes are set.
    """
    if not _is_indexable_html_page_url(url):
        return False
    path_value = _url_path_for_rules(url)
    for ex in exclude_rules:
        op, pat = ex.get("operator", ""), ex.get("pattern", "")
        if pat and _rule_matches(op, pat, path_value):
            return False
    if not include_rules:
        return True
    for inc in include_rules:
        op, pat = inc.get("operator", ""), inc.get("pattern", "")
        if pat and _rule_matches(op, pat, path_value):
            return True
    return False


def _url_passes_excludes_only(url: str, exclude_rules: list[dict[str, str]]) -> bool:
    """Used to decide if the crawl seed may be fetched for link discovery (include rules may still omit it)."""
    path_value = _url_path_for_rules(url)
    for ex in exclude_rules:
        op, pat = ex.get("operator", ""), ex.get("pattern", "")
        if pat and _rule_matches(op, pat, path_value):
            return False
    return True


def _filter_discovered_page_urls(
    urls: list[str],
    include_rules: list[dict[str, str]],
    exclude_rules: list[dict[str, str]],
) -> list[str]:
    """Apply path include/exclude to a full sitemap URL list (order preserved, deduped)."""
    out: list[str] = []
    seen: set[str] = set()
    for u in urls:
        if u in seen:
            continue
        if _url_passes_filters(u, include_rules, exclude_rules):
            seen.add(u)
            out.append(u)
    return out


def _cap_path_filtered_urls(
    urls: list[str],
    include_rules: list[dict[str, str]],
    exclude_rules: list[dict[str, str]],
    max_urls: int,
) -> list[str]:
    """Strict path filter before queueing HTML fetch (include = path must match)."""
    if not include_rules and not exclude_rules:
        return urls[:max_urls]
    filtered = _filter_discovered_page_urls(urls, include_rules, exclude_rules)
    return filtered[:max_urls]


def _extract_embedded_same_site_paths(base_url: str, html: str) -> list[str]:
    """Paths embedded in JSON/scripts (common on Shopify collection grids), not only ``<a href>``."""
    parsed = urlparse(base_url)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        return []
    origin = f"{parsed.scheme}://{parsed.netloc}"
    host = re.escape(parsed.netloc)
    out: list[str] = []
    path_re = re.compile(
        rf"(?:https?://{host})?"
        r"(/?(?:collections/[^\s\"'<>\\?#]+(?:/products/[^\s\"'<>\\?#]+)?|products/[^\s\"'<>\\?#]+))",
        re.IGNORECASE,
    )
    for m in path_re.finditer(html):
        raw_path = m.group(1)
        if not raw_path.startswith("/"):
            raw_path = f"/{raw_path}"
        try:
            out.append(_normalize_url_string(f"{origin}{raw_path}"))
        except AppError:
            continue
    abs_re = re.compile(rf"https?://{host}/[^\s\"'<>\\?#]+", re.IGNORECASE)
    for m in abs_re.finditer(html):
        try:
            candidate = _normalize_url_string(m.group(0))
        except AppError:
            continue
        if _same_site_host(parsed.netloc, urlparse(candidate).netloc):
            out.append(candidate)
    return out


def _extract_page_urls(base_url: str, html: str) -> list[str]:
    """Anchor links plus same-site paths embedded in page HTML."""
    seen: set[str] = set()
    merged: list[str] = []
    for u in _extract_links(base_url, html) + _extract_embedded_same_site_paths(base_url, html):
        if u in seen:
            continue
        seen.add(u)
        merged.append(u)
    return merged


async def _expand_urls_from_matching_hubs(
    seed_urls: list[str],
    *,
    include_rules: list[dict[str, str]],
    exclude_rules: list[dict[str, str]],
    max_urls: int,
) -> tuple[list[str], int]:
    """Fetch pages that matched include rules; add linked URLs only if their path still matches include/exclude."""
    if not include_rules or not seed_urls:
        return seed_urls[:max_urls], 0
    result: list[str] = []
    seen: set[str] = set()
    for u in seed_urls:
        if u not in seen:
            seen.add(u)
            result.append(u)
    added = 0
    async with httpx.AsyncClient(timeout=20, follow_redirects=True, headers=_WEBSITE_CRAWL_HEADERS) as client:
        for hub in seed_urls[:HUB_EXPAND_MAX_HUB_VISITS]:
            if len(result) >= max_urls:
                break
            try:
                response = await client.get(hub)
                response.raise_for_status()
                _charged, html = _charge_bytes_and_html_from_response(
                    response, remaining_budget_bytes=_CRAWL_BODY_CHARGE_CAP_BYTES
                )
                hub_added = 0
                for link in _extract_page_urls(hub, html):
                    if len(result) >= max_urls or hub_added >= HUB_EXPAND_MAX_LINKS_PER_HUB:
                        break
                    if link in seen:
                        continue
                    if not _url_passes_filters(link, include_rules, exclude_rules):
                        continue
                    seen.add(link)
                    result.append(link)
                    hub_added += 1
                    added += 1
            except Exception:
                continue
    return result[:max_urls], added


async def _require_knowledge_source_exists(db: AsyncSession, source_id: UUID) -> None:
    """Abort indexing if the source row was removed (e.g. user deleted it mid-crawl)."""
    row = (
        await db.execute(
            text("select 1 from public.knowledge_sources where id = :id limit 1"),
            {"id": str(source_id)},
        )
    ).first()
    if row is None:
        raise AppError(
            code="knowledge.source_removed",
            message="This website source was deleted while indexing was still running.",
            status_code=409,
        )


def _canonical_host(host: str) -> str:
    h = (host or "").strip().lower()
    if h.startswith("www."):
        return h[4:]
    return h


def _same_site_host(base_host: str, candidate_host: str) -> bool:
    return _canonical_host(base_host) == _canonical_host(candidate_host)


def _extract_links(base_url: str, html: str) -> list[str]:
    base = urlparse(base_url)
    soup = _parse_html(html)
    links: list[str] = []
    for anchor in soup.find_all("a", href=True):
        href = str(anchor["href"]).strip()
        if not href or href.startswith("#") or href.startswith("mailto:") or href.startswith("tel:"):
            continue
        absolute = urljoin(base_url, href)
        parsed = urlparse(absolute)
        if parsed.scheme not in {"http", "https"}:
            continue
        if not _same_site_host(base.netloc, parsed.netloc):
            continue
        normalized = f"{parsed.scheme}://{parsed.netloc}{parsed.path or '/'}"
        links.append(normalized)
    return links


# Meta / JSON keys that are layout, auth, or asset pointers — not useful in RAG excerpts.
_META_SKIP_NAMES: frozenset[str] = frozenset(
    {
        "viewport",
        "theme-color",
        "color-scheme",
        "robots",
        "referrer",
        "format-detection",
        "msapplication-tilecolor",
        "csrf-token",
        "csrf-param",
        "google-site-verification",
        "facebook-domain-verification",
        "p:domain_verify",
    }
)
_META_SKIP_PROPERTY_PREFIXES: tuple[str, ...] = (
    "og:image",
    "og:video",
    "twitter:image",
    "twitter:player",
)

_JSON_SKIP_KEYS: frozenset[str] = frozenset(
    {
        "@context",
        "context",
        "@id",
        "accessToken",
        "betas",
        "shopId",
        "predictiveSearch",
        "locale",
        "domain",
        "featured_image",
        "image",
        "images",
        "logo",
        "sameAs",
        "seller",
        "hasMerchantReturnPolicy",
        "shippingDetails",
        "inventory_management",
        "selling_plan_allocations",
        "quantity_rule",
        "requires_selling_plan",
    }
)
_JSON_MONEY_KEYS: frozenset[str] = frozenset(
    {
        "price",
        "compare_at_price",
        "minprice",
        "maxprice",
        "lowprice",
        "highprice",
        "unitprice",
    }
)
_MAX_JSON_INDEX_FRAGMENTS = 200
_THIN_PAGE_BODY_FALLBACK_CHARS = 700


def _humanize_json_key(key: str) -> str:
    clean = key.lstrip("@").replace("_", " ").replace("-", " ")
    if not clean:
        return "value"
    return clean[0].upper() + clean[1:] if len(clean) > 1 else clean.upper()


def _json_key_is_noise(key: str) -> bool:
    lowered = key.lower().lstrip("@")
    if lowered in _JSON_SKIP_KEYS or lowered == "type":
        return True
    if lowered in {"id", "gid", "shopid"} or lowered.endswith("id") and lowered not in {
        "sku",
        "grid",
    }:
        return True
    if "token" in lowered or "secret" in lowered or "password" in lowered:
        return True
    if lowered.endswith("url") or lowered.endswith("uri") or lowered in {"url", "href", "src"}:
        return True
    return False


def _scalar_value_indexable(value: object) -> bool:
    if isinstance(value, bool):
        return True
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return True
    if isinstance(value, str):
        text = _normalize_text(value)
        if len(text) < 2:
            return False
        if text.startswith(("http://", "https://", "//")):
            return False
        if len(text) > 800:
            return False
        return True
    return False


def _parse_money_amount(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        numeric = float(value)
        if isinstance(value, int) and value >= 10_000 and value % 100 == 0:
            return value / 100.0
        return numeric
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if not text:
            return None
        try:
            return float(text)
        except ValueError:
            return None
    return None


def _format_json_scalar(key: str, value: object, parent: dict[str, object] | None) -> str | None:
    key_l = key.lower().lstrip("@")
    if isinstance(value, bool):
        return f"{_humanize_json_key(key)}: {'yes' if value else 'no'}"
    if key_l in _JSON_MONEY_KEYS:
        numeric = _parse_money_amount(value)
        if numeric is not None:
            currency = None
            if parent:
                raw_currency = parent.get("priceCurrency") or parent.get("currency")
                if isinstance(raw_currency, str) and raw_currency.strip():
                    currency = raw_currency.strip()
            amount = (
                str(int(numeric))
                if numeric == int(numeric)
                else f"{numeric:.2f}".rstrip("0").rstrip(".")
            )
            if currency:
                return f"{currency} {amount}"
            return f"{_humanize_json_key(key)}: {amount}"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f"{_humanize_json_key(key)}: {value}"
    if isinstance(value, str):
        text = _normalize_text(value)
        if not text:
            return None
        return f"{_humanize_json_key(key)}: {text}"
    return None


def _flatten_json_for_indexing(
    node: object,
    fragments: list[str],
    *,
    depth: int = 0,
    max_depth: int = 14,
    parent: dict[str, object] | None = None,
) -> None:
    """Recursively turn JSON-LD / inline JSON into labeled plain-text snippets."""
    if depth > max_depth or len(fragments) >= _MAX_JSON_INDEX_FRAGMENTS:
        return
    if isinstance(node, dict):
        price = node.get("price")
        currency = node.get("priceCurrency") or node.get("currency")
        if price is not None and isinstance(currency, str) and currency.strip():
            combined = _format_json_scalar("price", price, {**node, "priceCurrency": currency})
            if combined and combined not in fragments:
                fragments.append(combined)
        for key, value in node.items():
            if _json_key_is_noise(key):
                continue
            if key in ("price", "priceCurrency", "currency") and price is not None and currency:
                continue
            if isinstance(value, dict):
                _flatten_json_for_indexing(value, fragments, depth=depth + 1, max_depth=max_depth, parent=node)
            elif isinstance(value, list):
                _flatten_json_for_indexing(value, fragments, depth=depth + 1, max_depth=max_depth, parent=node)
            elif _scalar_value_indexable(value):
                line = _format_json_scalar(key, value, node)
                if line:
                    fragments.append(line)
    elif isinstance(node, list):
        primitive_items = [item for item in node if _scalar_value_indexable(item)]
        if primitive_items and len(primitive_items) == len(node) and parent is not None:
            joined = ", ".join(_normalize_text(str(item)) for item in primitive_items)
            parent_key = ""
            if isinstance(parent, dict):
                for k, v in parent.items():
                    if v is node:
                        parent_key = k
                        break
            label = _humanize_json_key(parent_key) if parent_key else "Items"
            fragments.append(f"{label}: {joined}")
            return
        for item in node:
            _flatten_json_for_indexing(item, fragments, depth=depth + 1, max_depth=max_depth, parent=parent)


def _head_meta_text_fragments(soup: BeautifulSoup) -> list[str]:
    """
    Visible text from ``get_text`` omits ``<meta content=\"...\">`` — collect all
    non-layout meta (Open Graph, Twitter, product tags, descriptions).
    """
    fragments: list[str] = []
    title = soup.find("title")
    if title:
        t = _normalize_text(title.get_text(" "))
        if t:
            fragments.append(t)
    for meta in soup.find_all("meta"):
        content = meta.get("content")
        if not content or not str(content).strip():
            continue
        name = str(meta.get("name") or "").strip()
        prop = str(meta.get("property") or "").strip()
        label = prop or name
        if not label:
            continue
        label_l = label.lower()
        if label_l in _META_SKIP_NAMES:
            continue
        if any(label_l.startswith(prefix) for prefix in _META_SKIP_PROPERTY_PREFIXES):
            continue
        if label_l.startswith("og:image:") or label_l.startswith("twitter:image:"):
            continue
        text = _normalize_text(str(content))
        if not text:
            continue
        fragments.append(f"{label}: {text}")
    return fragments


def _microdata_text_fragments(soup: BeautifulSoup) -> list[str]:
    fragments: list[str] = []
    for el in soup.find_all(attrs={"itemprop": True}):
        prop = str(el.get("itemprop") or "").strip()
        if not prop or _json_key_is_noise(prop):
            continue
        if el.name == "meta":
            value = el.get("content")
        elif el.name in {"a", "link", "img", "source"}:
            value = el.get("href") or el.get("src") or el.get("content")
        else:
            value = el.get("content") or el.get_text(" ", strip=True)
        if value is None:
            continue
        line = _format_json_scalar(prop, value, None)
        if line:
            fragments.append(line)
    return fragments


def _embedded_json_text_fragments(soup: BeautifulSoup) -> list[str]:
    """JSON-LD and inline ``application/json`` blocks (Shopify variants, FAQs, etc.)."""
    fragments: list[str] = []
    for script in soup.find_all("script"):
        script_type = str(script.get("type") or "").lower()
        if script_type not in {"application/ld+json", "application/json"}:
            continue
        raw = script.string or script.get_text(" ", strip=True)
        if not raw or len(raw.strip()) < 24:
            continue
        try:
            payload = json.loads(raw)
        except json.JSONDecodeError:
            continue
        _flatten_json_for_indexing(payload, fragments)
    return fragments


def _dedupe_preserve_order_snippets(snippets: list[str]) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for s in snippets:
        n = _normalize_text(s)
        if len(n) < 2:
            continue
        key = n.casefold()
        if key in seen:
            continue
        seen.add(key)
        out.append(n)
    return out


def _extract_structural_body_blocks(root: BeautifulSoup) -> list[str]:
    structural_blocks: list[str] = []
    block_tags = [
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "p",
        "li",
        "dt",
        "dd",
        "pre",
        "code",
        "table",
        "blockquote",
        "figcaption",
    ]
    for tag in root.find_all(block_tags):
        name = str(tag.name).lower()
        text_val = _normalize_text(tag.get_text(" ", strip=True))
        if not text_val:
            continue
        if name.startswith("h"):
            level = int(name[1]) if len(name) > 1 and name[1].isdigit() else 2
            structural_blocks.append(f"{'#' * max(1, min(level, 6))} {text_val}")
        elif name in {"pre", "code"}:
            structural_blocks.append(f"```\n{text_val}\n```")
        elif name == "table":
            structural_blocks.append(f"[Table]\n{text_val}")
        else:
            structural_blocks.append(text_val)
    return structural_blocks


def _extract_page_text(html: str) -> str:
    soup = _parse_html(html)
    head_frags = _head_meta_text_fragments(soup)
    json_frags = _embedded_json_text_fragments(soup)
    microdata_frags = _microdata_text_fragments(soup)
    for element in soup(["script", "style", "noscript"]):
        element.decompose()

    main_candidates = soup.select("main, article, [role='main'], #MainContent, #main-content")
    root = main_candidates[0] if main_candidates else (soup.body or soup)

    structural_blocks = _extract_structural_body_blocks(root)
    merged = _dedupe_preserve_order_snippets(
        [*head_frags, *json_frags, *microdata_frags, *structural_blocks]
    )
    if len("\n\n".join(merged)) < _THIN_PAGE_BODY_FALLBACK_CHARS:
        fallback = _normalize_text(root.get_text("\n", strip=True))
        if fallback:
            merged = _dedupe_preserve_order_snippets([*merged, fallback])
    return "\n\n".join(merged).strip()


def _extract_page_title(html: str) -> str | None:
    soup = _parse_html(html)
    title = soup.find("title")
    if not title:
        return None
    text_value = _normalize_text(title.get_text(" "))
    return text_value or None


def _extract_social_preview_image(html: str, page_url: str) -> str | None:
    soup = _parse_html(html)
    raw_candidates: list[str] = []
    for tag in soup.find_all("meta"):
        prop = tag.get("property")
        if prop and str(prop).lower() in ("og:image", "og:image:url", "og:image:secure_url"):
            c = tag.get("content")
            if c and str(c).strip():
                raw_candidates.append(str(c).strip())
        name = tag.get("name")
        if name and str(name).lower() in ("twitter:image", "twitter:image:src"):
            c = tag.get("content")
            if c and str(c).strip():
                raw_candidates.append(str(c).strip())
    for tag in soup.find_all("link", href=True):
        rel = tag.get("rel")
        rel_parts = rel if isinstance(rel, list) else ([rel] if rel else [])
        if any(str(r).lower() == "image_src" for r in rel_parts):
            href = tag.get("href")
            if href and str(href).strip():
                raw_candidates.append(str(href).strip())
    for raw in raw_candidates:
        absolute = urljoin(page_url, raw)
        parsed = urlparse(absolute)
        if parsed.scheme in {"http", "https"} and parsed.netloc:
            return absolute
    return None


def _embedding_api_error_message(status_code: int, body: str) -> str:
    try:
        payload = json.loads(body)
        err = payload.get("error") if isinstance(payload, dict) else None
        if isinstance(err, dict) and err.get("message"):
            return str(err["message"])[:500]
    except json.JSONDecodeError:
        pass
    return f"HTTP {status_code}"


async def _post_one_embedding_batch(
    client: httpx.AsyncClient,
    batch: list[str],
    *,
    allow_split: bool,
    settings: Any,
) -> tuple[list[list[float]], int]:
    """One OpenAI embeddings HTTP call; returns vectors and API-reported prompt/total tokens."""
    response = await client.post(
        "https://api.openai.com/v1/embeddings",
        headers={"Authorization": f"Bearer {settings.openai_api_key}"},
        json={"model": settings.openai_embedding_model, "input": batch},
    )
    if response.status_code >= 400:
        body = response.text or ""
        lower = body.lower()
        oversize = (
            response.status_code == 400
            and allow_split
            and len(batch) > 1
            and (
                "maximum request size" in lower
                or "too many tokens" in lower
                or "context length" in lower
            )
        )
        if oversize:
            mid = len(batch) // 2
            left_v, left_t = await _post_one_embedding_batch(
                client, batch[:mid], allow_split=True, settings=settings
            )
            right_v, right_t = await _post_one_embedding_batch(
                client, batch[mid:], allow_split=True, settings=settings
            )
            return left_v + right_v, left_t + right_t
        log.warning(
            "embedding_batch_failed",
            status_code=response.status_code,
            body_preview=body[:500],
            batch_inputs=len(batch),
        )
        detail_msg = _embedding_api_error_message(response.status_code, body)
        raise AppError(
            code="knowledge.embedding_failed",
            message=f"Indexing API error: {detail_msg}",
            status_code=502,
            details={"status_code": response.status_code, "body": body[:800]},
        )
    payload = response.json()
    vectors = [item["embedding"] for item in payload.get("data", [])]
    if len(vectors) != len(batch):
        raise AppError(code="knowledge.embedding_failed", message="Indexing response length mismatch", status_code=502)
    usage = payload.get("usage") if isinstance(payload.get("usage"), dict) else {}
    tokens = int(usage.get("prompt_tokens") or usage.get("total_tokens") or 0)
    return vectors, tokens


async def _embed_texts(chunks: list[str]) -> list[list[float]]:
    settings = get_settings()
    if not settings.openai_api_key:
        raise AppError(
            code="knowledge.embedding_not_configured",
            message="Knowledge indexing is not configured on the server",
            status_code=500,
        )

    if not chunks:
        return []

    all_vectors: list[list[float]] = []
    idx = 0
    async with httpx.AsyncClient(timeout=120) as client:
        while idx < len(chunks):
            batch: list[str] = []
            batch_tokens = 0
            while idx < len(chunks) and len(batch) < _EMBED_BATCH_MAX_INPUTS:
                next_tok = _approx_embed_request_tokens(chunks[idx])
                if batch and batch_tokens + next_tok > _EMBED_BATCH_MAX_TOKENS_EST:
                    break
                batch.append(chunks[idx])
                batch_tokens += next_tok
                idx += 1
            if not batch:
                batch = [chunks[idx]]
                idx += 1

            vecs, _tokens = await _post_one_embedding_batch(
                client, batch, allow_split=True, settings=settings
            )
            all_vectors.extend(vecs)

    if len(all_vectors) != len(chunks):
        raise AppError(code="knowledge.embedding_failed", message="Indexing response length mismatch", status_code=502)
    return all_vectors


async def embed_texts_with_token_usage(chunks: list[str]) -> tuple[list[list[float]], int]:
    """Embeddings with summed OpenAI usage tokens across all HTTP batches (RAG / runtime billing)."""
    settings = get_settings()
    if not settings.openai_api_key:
        raise AppError(
            code="knowledge.embedding_not_configured",
            message="Knowledge indexing is not configured on the server",
            status_code=500,
        )
    if not chunks:
        return [], 0

    all_vectors: list[list[float]] = []
    total_tokens = 0
    idx = 0
    async with httpx.AsyncClient(timeout=120) as client:
        while idx < len(chunks):
            batch: list[str] = []
            batch_tokens = 0
            while idx < len(chunks) and len(batch) < _EMBED_BATCH_MAX_INPUTS:
                next_tok = _approx_embed_request_tokens(chunks[idx])
                if batch and batch_tokens + next_tok > _EMBED_BATCH_MAX_TOKENS_EST:
                    break
                batch.append(chunks[idx])
                batch_tokens += next_tok
                idx += 1
            if not batch:
                batch = [chunks[idx]]
                idx += 1

            vecs, tok = await _post_one_embedding_batch(client, batch, allow_split=True, settings=settings)
            all_vectors.extend(vecs)
            total_tokens += tok

    if len(all_vectors) != len(chunks):
        raise AppError(code="knowledge.embedding_failed", message="Indexing response length mismatch", status_code=502)
    return all_vectors, total_tokens


def _vector_literal(values: list[float]) -> str:
    return "[" + ",".join(f"{x:.10f}" for x in values) + "]"


def _validate_source_payload(payload: KnowledgeSourceCreateRequest) -> None:
    if payload.type == "website" and not payload.source_url:
        raise AppError(code="validation.invalid_input", message="source_url is required for website type", status_code=422)
    if payload.type == "file" and (not payload.storage_bucket or not payload.storage_path):
        raise AppError(
            code="validation.invalid_input",
            message="storage_bucket and storage_path are required for file type",
            status_code=422,
        )
    if payload.type == "text_snippet":
        body = (payload.raw_text or "").strip()
        if not body:
            raise AppError(code="validation.invalid_input", message="raw_text is required for text_snippet type", status_code=422)


async def create_source(db: AsyncSession, user_id: UUID, payload: KnowledgeSourceCreateRequest) -> KnowledgeSourceDTO:
    _validate_source_payload(payload)

    src_status = (payload.status or "pending").strip() or "pending"
    raw_for_insert = (payload.raw_text or "").strip() if payload.type == "text_snippet" else None

    result = await db.execute(
        text(
            """
            insert into public.knowledge_sources (
              agent_id, user_id, type, title, status, source_url, storage_bucket, storage_path, raw_text, metadata
            ) values (
              :agent_id, :user_id, :type, :title, CAST(:status AS knowledge_source_status), :source_url, :storage_bucket, :storage_path, :raw_text, CAST(:metadata AS jsonb)
            )
            returning
              id, agent_id, user_id, type, title, status, source_url, storage_bucket, storage_path, raw_text,
              metadata, error_message, last_indexed_at, created_at, updated_at
            """
        ),
        {
            "agent_id": str(payload.agent_id),
            "user_id": str(user_id),
            "type": payload.type,
            "title": payload.title,
            "status": src_status,
            "source_url": payload.source_url,
            "storage_bucket": payload.storage_bucket,
            "storage_path": payload.storage_path,
            "raw_text": raw_for_insert,
            "metadata": json.dumps(payload.metadata or {}),
        },
    )
    await db.commit()
    return KnowledgeSourceDTO.model_validate(result.mappings().one())


async def list_sources(db: AsyncSession, user_id: UUID, agent_id: UUID | None = None) -> list[KnowledgeSourceDTO]:
    sql = """
        select
          id, agent_id, user_id, type, title, status, source_url, storage_bucket, storage_path, raw_text,
          metadata, error_message, last_indexed_at, created_at, updated_at
        from public.knowledge_sources
        where user_id = :user_id
    """
    params: dict[str, object] = {"user_id": str(user_id)}
    if agent_id:
        sql += " and agent_id = :agent_id"
        params["agent_id"] = str(agent_id)
    sql += " order by created_at desc"
    result = await db.execute(text(sql), params)
    return [KnowledgeSourceDTO.model_validate(row) for row in result.mappings().all()]


async def _create_job(db: AsyncSession, source: KnowledgeSourceDTO, user_id: UUID) -> IndexJobDTO:
    result = await db.execute(
        text(
            """
            insert into public.indexing_jobs (
              knowledge_source_id, agent_id, user_id, status, triggered_by, phase, progress_pct
            ) values (
              :knowledge_source_id, :agent_id, :user_id, 'queued', 'api', 'queued', 0
            )
            returning
              id, knowledge_source_id, agent_id, user_id, status, attempt, triggered_by, error_message,
              started_at, finished_at, phase, pages_total, pages_processed, chunks_total, chunks_embedded, progress_pct, metrics, created_at, updated_at
            """
        ),
        {
            "knowledge_source_id": str(source.id),
            "agent_id": str(source.agent_id),
            "user_id": str(user_id),
        },
    )
    return IndexJobDTO.model_validate(result.mappings().one())


async def _set_job_running(db: AsyncSession, job_id: UUID) -> None:
    await db.execute(
        text(
            """
            update public.indexing_jobs
            set status = 'running', phase = 'crawling', started_at = coalesce(started_at, now()), error_message = null, progress_pct = 5
            where id = :job_id
            """
        ),
        {"job_id": str(job_id)},
    )


async def _load_source(db: AsyncSession, source_id: UUID, user_id: UUID) -> KnowledgeSourceDTO:
    result = await db.execute(
        text(
            """
            select
              id, agent_id, user_id, type, title, status, source_url, storage_bucket, storage_path, raw_text,
              metadata, error_message, last_indexed_at, created_at, updated_at
            from public.knowledge_sources
            where id = :source_id and user_id = :user_id
            """
        ),
        {"source_id": str(source_id), "user_id": str(user_id)},
    )
    row = result.mappings().first()
    if row is None:
        raise AppError(code="knowledge.source_not_found", message="Knowledge source not found", status_code=404)
    return KnowledgeSourceDTO.model_validate(row)


async def enqueue_index_website_source(
    db: AsyncSession, source_id: UUID, user_id: UUID, *, mark_running: bool = True
) -> tuple[KnowledgeSourceDTO, IndexJobDTO]:
    source = await _load_source(db, source_id, user_id)
    if source.type != "website":
        raise AppError(code="validation.invalid_input", message="Only website sources are indexable in this step", status_code=422)
    if not source.source_url:
        raise AppError(code="validation.invalid_input", message="Website source URL is missing", status_code=422)

    job = await _create_job(db, source, user_id)
    if mark_running:
        await _set_job_running(db, job.id)
    await db.execute(
        text(
            """
            update public.knowledge_sources
            set status = 'indexing', error_message = null
            where id = :source_id
            """
        ),
        {"source_id": str(source.id)},
    )
    await db.commit()
    return source, job


async def enqueue_index_website_source_queued(
    db: AsyncSession, source_id: UUID, user_id: UUID
) -> tuple[KnowledgeSourceDTO, IndexJobDTO]:
    """Leave the job `queued` for `indexing_worker` (dashboard website crawl uses this path)."""
    return await enqueue_index_website_source(db, source_id, user_id, mark_running=False)


def _rules_from_payload(rules: list[WebsitePathRule]) -> list[dict[str, str]]:
    """Persist only rules with a non-empty pattern; blank chips would otherwise break includes entirely.

    For ``starts_with`` and ``exact_match``, path rules match URL path segments (``/products/...``);
    if the user omits a leading ``/``, add it so ``products`` matches ``/products``.
    """
    out: list[dict[str, str]] = []
    for r in rules:
        pat = (r.pattern or "").strip()
        if not pat:
            continue
        if r.operator in ("starts_with", "exact_match") and not pat.startswith("/"):
            pat = f"/{pat.lstrip('/')}"
        out.append({"operator": r.operator, "pattern": pat})
    return out


async def create_and_enqueue_dashboard_website(
    db: AsyncSession,
    user_id: UUID,
    payload: WebsiteIngestBase | WebsiteIndividualRequest,
    mode: WebsiteMode,
) -> tuple[KnowledgeSourceDTO, IndexJobDTO | None]:
    if isinstance(payload, WebsiteIndividualRequest):
        source_url = normalize_dashboard_website_url(payload.protocol, payload.url_input)
        title = payload.title or urlparse(source_url).netloc or "Website"
        include_rules: list[WebsitePathRule] = []
        exclude_rules: list[WebsitePathRule] = []
    else:
        source_url = normalize_dashboard_website_url(payload.protocol, payload.url_input)
        title = payload.title or urlparse(source_url).netloc or "Website"
        include_rules = payload.include_rules
        exclude_rules = payload.exclude_rules

    metadata: dict[str, object] = {
        "origin": "dashboard_website",
        "website_mode": mode,
        "include_rules": _rules_from_payload(include_rules),
        "exclude_rules": _rules_from_payload(exclude_rules),
        "max_pages": MAX_DASHBOARD_WEBSITE_PAGES,
    }

    dup = await _find_dashboard_website_duplicate(
        db,
        user_id,
        payload.agent_id,
        source_url,
        include_rules=_rules_from_payload(include_rules),
        exclude_rules=_rules_from_payload(exclude_rules),
    )
    if dup is not None and mode != "individual":
        # Don't block a new dashboard crawl just because some seed/landing page was indexed
        # previously (e.g. onboarding preview). We dedupe at per-URL embedding time instead.
        existing_id, dup_reason = dup
        metadata = {
            **metadata,
            "duplicate_of_source_id": str(existing_id),
            "duplicate_reason": dup_reason,
        }
    if dup is not None and mode == "individual":
        existing_id, dup_reason = dup
        metadata = {
            **metadata,
            "reindexed_duplicate": True,
            "duplicate_of_source_id": str(existing_id),
            "duplicate_reason": dup_reason,
        }

    source = await create_source(
        db,
        user_id,
        KnowledgeSourceCreateRequest(
            agent_id=payload.agent_id,
            type="website",
            title=title,
            source_url=source_url,
            metadata=metadata,
        ),
    )
    _, job = await enqueue_index_website_source_queued(db, source.id, user_id)
    refreshed_source = await _load_source(db, source.id, user_id)
    refreshed_job = await get_latest_job(db, source.id, user_id)
    return refreshed_source, refreshed_job or job


def _app_error_for_empty_sitemap_discovery(d: SitemapDiscoveryResult) -> AppError:
    """Use when ``d.urls`` is empty for explicit sitemap ingest (not dashboard crawl fallback)."""
    if d.http_success_count == 0 and d.http_failure_count > 0:
        return AppError(
            code="knowledge.sitemap_unreachable",
            message=(
                "Could not fetch sitemap XML (HTTP error, TLS, or blocked). "
                "Use a direct https://…/sitemap.xml URL."
            ),
            status_code=422,
        )
    if d.http_success_count > 0 and not d.had_successful_xml_document:
        return AppError(
            code="knowledge.sitemap_invalid",
            message=(
                "The response was not valid sitemap XML (often HTML). "
                "Open Graph preview URLs must point to an XML sitemap, not the storefront homepage."
            ),
            status_code=422,
        )
    return AppError(
        code="knowledge.sitemap_empty",
        message=(
            "No page URLs in the sitemap matched your path filters, or the sitemap lists no pages. "
            "Loosen include/exclude rules or point to a different sitemap."
        ),
        status_code=422,
    )


async def preview_dashboard_website_filtered_urls(
    payload: WebsiteUrlPreviewRequest,
) -> WebsiteUrlPreviewResponse:
    """Return filtered sitemap URL count and a sample (no source row, no indexing)."""
    source_url = normalize_dashboard_website_url(payload.protocol, payload.url_input)
    inc = _rules_from_payload(payload.include_rules)
    exc = _rules_from_payload(payload.exclude_rules)
    discovered = await _discover_sitemap_urls(source_url, MAX_DASHBOARD_WEBSITE_PAGES, inc, exc)
    urls = discovered.urls
    cap = min(int(payload.max_sample_urls), 100)
    warning_parts: list[str] = []
    if discovered.truncated_by_doc_cap:
        warning_parts.append(
            "Sitemap discovery stopped after many nested sitemap files; the URL list may be incomplete."
        )
    if urls and discovered.http_failure_count > 0:
        warning_parts.append(
            "Some nested sitemaps failed to load; the URL count may be incomplete."
        )
    if discovered.urls_discovered_total > 0 and (inc or exc):
        filter_note = (
            f"Found {discovered.urls_discovered_total} page URL(s) in sitemap XML; "
            f"{discovered.urls_after_filter} match your path filters"
        )
        if discovered.hub_urls_added > 0:
            filter_note += f"; +{discovered.hub_urls_added} from matching hub pages (path still matches)"
        filter_note += f"; {len(urls)} will be indexed (cap {MAX_DASHBOARD_WEBSITE_PAGES})."
        warning_parts.append(filter_note)
    warning = " ".join(warning_parts) if warning_parts else None
    if urls:
        return WebsiteUrlPreviewResponse(
            discovery_mode="sitemap",
            filtered_url_count=len(urls),
            sample_urls=urls[:cap],
            truncated=len(urls) > cap,
            message=None,
            discovery_warning=warning,
            sitemap_truncated=discovered.truncated_by_doc_cap,
        )
    if discovered.http_success_count == 0 and discovered.http_failure_count > 0:
        return WebsiteUrlPreviewResponse(
            discovery_mode="sitemap_unreachable",
            filtered_url_count=0,
            sample_urls=[],
            truncated=False,
            message=(
                "Could not fetch sitemap XML (HTTP error, TLS, or blocked). "
                "Use a direct https://…/sitemap.xml URL."
            ),
            discovery_warning=warning,
            sitemap_truncated=discovered.truncated_by_doc_cap,
        )
    if discovered.http_success_count > 0 and not discovered.had_successful_xml_document:
        return WebsiteUrlPreviewResponse(
            discovery_mode="sitemap_invalid",
            filtered_url_count=0,
            sample_urls=[],
            truncated=False,
            message=(
                "The response was not valid sitemap XML (often HTML). Point to sitemap.xml, not the homepage."
            ),
            discovery_warning=warning,
            sitemap_truncated=discovered.truncated_by_doc_cap,
        )
    return WebsiteUrlPreviewResponse(
        discovery_mode="bfs_fallback",
        filtered_url_count=0,
        sample_urls=[],
        truncated=False,
        message=(
            "No matching URLs found in sitemap XML for this site and filters. "
            "A full crawl would use link following (BFS); URL count is not known until the crawl runs."
        ),
        discovery_warning=warning,
        sitemap_truncated=discovered.truncated_by_doc_cap,
    )


async def _create_crawl_run(db: AsyncSession, source: KnowledgeSourceDTO, user_id: UUID, settings: dict[str, object]) -> UUID:
    result = await db.execute(
        text(
            """
            insert into public.knowledge_crawl_runs (knowledge_source_id, agent_id, user_id, status, started_at, settings)
            values (:source_id, :agent_id, :user_id, 'running', now(), cast(:settings as jsonb))
            returning id
            """
        ),
        {
            "source_id": str(source.id),
            "agent_id": str(source.agent_id),
            "user_id": str(user_id),
            "settings": json.dumps(settings),
        },
    )
    return UUID(str(result.mappings().one()["id"]))


def _ingest_params_from_metadata(metadata: dict[str, object]) -> tuple[WebsiteMode, int, list[dict[str, str]], list[dict[str, str]]]:
    origin = str(metadata.get("origin") or "")
    mode_raw = metadata.get("website_mode")
    if mode_raw in ("crawl", "sitemap", "individual"):
        mode: WebsiteMode = mode_raw  # type: ignore[assignment]
    elif origin == "dashboard_website":
        mode = "crawl"
    else:
        mode = "crawl"
    max_pages = int(metadata.get("max_pages") or MAX_DASHBOARD_WEBSITE_PAGES)
    include = [dict(x) for x in (metadata.get("include_rules") or []) if isinstance(x, dict)]
    exclude = [dict(x) for x in (metadata.get("exclude_rules") or []) if isinstance(x, dict)]
    return mode, max_pages, include, exclude


def _website_ingest_uses_live_discovery(metadata: dict[str, object]) -> bool:
    """Dashboard and onboarding: sitemap seed + incremental page rows while the worker crawls."""
    return str(metadata.get("origin") or "") in ("dashboard_website", "onboarding")


async def _crawl_pages(
    seed_url: str,
    max_pages: int,
    include_rules: list[dict[str, str]],
    exclude_rules: list[dict[str, str]],
    *,
    crawl_budget_bytes: int,
    progress_hook: Callable[..., Awaitable[None]] | None = None,
    progress_every: int = DASHBOARD_CRAWL_PROGRESS_EVERY,
) -> tuple[list[dict[str, object]], int, str | None, int, str, int]:
    """Crawl HTML pages; stop when cumulative extracted text bytes exceed ``crawl_budget_bytes``.

    The seed URL is always fetched if it passes **exclude** rules, even when **include** rules
    omit it (e.g. homepage ``/`` while include only paths matching ``/tapered/``). Those visits
    are used for link discovery; only URLs that pass the full include/exclude filters are indexed.

    When includes are non-empty, outbound links are queued using **exclude** rules only (plus hop /
    visit caps) so category/hub pages can sit on paths that do not yet match the include pattern.
    """
    parsed_seed = urlparse(seed_url)
    seed_normalized = f"{parsed_seed.scheme}://{parsed_seed.netloc}{parsed_seed.path or '/'}"
    if not _url_passes_excludes_only(seed_normalized, exclude_rules):
        return [], 0, None, 0, "complete", 0
    if not include_rules and not _url_passes_filters(seed_normalized, include_rules, exclude_rules):
        return [], 0, None, 0, "complete", 0

    queue: deque[tuple[str, int]] = deque([(seed_normalized, 0)])
    seen: set[str] = set()
    pages: list[dict[str, object]] = []
    links_discovered = 0
    preview_image_url: str | None = None
    stored_text_bytes = 0
    stopped_reason = "complete"
    http_visits = 0

    async with httpx.AsyncClient(timeout=20, follow_redirects=True, headers=_WEBSITE_CRAWL_HEADERS) as client:
        while queue and len(pages) < max_pages:
            if stored_text_bytes >= crawl_budget_bytes and len(pages) > 0:
                stopped_reason = "budget"
                break
            url, depth = queue.popleft()
            if url in seen:
                continue
            seen.add(url)
            if include_rules and len(seen) > DASHBOARD_BFS_INCLUDE_RULE_MAX_PAGES_VISITED:
                stopped_reason = "discovery_cap"
                break
            try:
                http_visits += 1
                response = await client.get(url)
                status_code = int(response.status_code)
                response.raise_for_status()
                _charged, html = _charge_bytes_and_html_from_response(
                    response, remaining_budget_bytes=_CRAWL_BODY_CHARGE_CAP_BYTES
                )
                if preview_image_url is None:
                    preview_image_url = _extract_social_preview_image(html, url)
                page_text = _extract_page_text(html)
                page_title = _extract_page_title(html)
                links = _extract_links(url, html)
                links_discovered += len(links)
                allow_more_links = stored_text_bytes < crawl_budget_bytes
                for link in links:
                    if link in seen:
                        continue
                    if include_rules:
                        if depth + 1 > DASHBOARD_BFS_INCLUDE_RULE_MAX_HOPS:
                            continue
                        passes_next_hop = _url_passes_excludes_only(link, exclude_rules)
                    else:
                        passes_next_hop = _url_passes_filters(link, include_rules, exclude_rules)
                    if not passes_next_hop:
                        continue
                    if allow_more_links and len(pages) + len(queue) < max_pages * 4:
                        queue.append((link, depth + 1))
                if _url_passes_filters(url, include_rules, exclude_rules):
                    remaining_storage = max(0, crawl_budget_bytes - stored_text_bytes)
                    if remaining_storage <= 0 and len(pages) > 0:
                        stopped_reason = "budget"
                        break
                    page_text = _truncate_utf8_to_bytes(page_text, remaining_storage)
                    text_bytes = len(page_text.encode("utf-8")) if page_text else 0
                    stored_text_bytes += text_bytes
                    pages.append(
                        {
                            "url": url,
                            "depth": depth,
                            "http_status": status_code,
                            "text": page_text,
                            "title": page_title,
                        }
                    )
                    if progress_hook and (
                        len(pages) % max(1, progress_every) == 0
                        or len(pages) >= max_pages
                        or http_visits % 8 == 0
                    ):
                        await progress_hook(list(pages), stored_text_bytes, http_visits)
                    if stored_text_bytes >= crawl_budget_bytes:
                        stopped_reason = "budget"
                        break
                    if len(pages) >= max_pages:
                        stopped_reason = "max_pages_safety"
                        break
            except Exception:
                if _url_passes_filters(url, include_rules, exclude_rules):
                    pages.append({"url": url, "depth": depth, "http_status": None, "text": "", "title": None})
                    if progress_hook and len(pages) % max(1, progress_every) == 0:
                        await progress_hook(list(pages), stored_text_bytes, http_visits)
    if stopped_reason == "complete" and len(pages) < max_pages:
        stopped_reason = "no_more_links"
    return pages[:max_pages], links_discovered, preview_image_url, stored_text_bytes, stopped_reason, http_visits


def _sitemap_seed_urls(start_url: str) -> list[str]:
    """URLs to treat as sitemap documents. Site-root URLs also try /sitemap.xml (homepage is usually HTML)."""
    try:
        norm = _normalize_url_string(start_url)
    except AppError:
        norm = start_url.strip()
    seeds = [norm]
    parsed = urlparse(norm)
    if parsed.scheme in ("http", "https") and parsed.netloc and (parsed.path or "/") == "/":
        candidate = f"{parsed.scheme}://{parsed.netloc}/sitemap.xml"
        if candidate != norm:
            seeds.append(candidate)
    return seeds


async def _discover_sitemap_urls(
    start_url: str,
    max_urls: int,
    include_rules: list[dict[str, str]],
    exclude_rules: list[dict[str, str]],
    *,
    progress: Callable[[int, int], Awaitable[None]] | None = None,
) -> SitemapDiscoveryResult:
    """Discover page URLs from sitemap XML (supports nested index files).

    Walks every nested sitemap (path filters are **not** applied during XML parse). After the
    full list is collected, include/exclude rules run on each page URL path. When include rules
    are set, matching hub pages (e.g. filtered collections) are fetched once to add embedded
    same-site URLs (e.g. ``/collections/…/products/…``) that still pass the same path filters.
    """
    all_page_urls: list[str] = []
    seen_page_urls: set[str] = set()
    seen_sitemaps: set[str] = set()
    sitemap_queue: deque[tuple[str, int]] = deque((u, 0) for u in _sitemap_seed_urls(start_url))
    documents_fetched = 0
    http_success_count = 0
    http_failure_count = 0
    xml_parse_failure_count = 0
    had_successful_xml_document = False
    truncated_by_doc_cap = False
    truncated_by_url_collect_cap = False

    # Large regional catalog XML can exceed default read timeouts; discovery only uses this client.
    sitemap_timeout = httpx.Timeout(180.0, connect=20.0)
    async with httpx.AsyncClient(timeout=sitemap_timeout, follow_redirects=True, headers=_WEBSITE_CRAWL_HEADERS) as client:

        async def _fetch_sitemap_body(sm_url: str) -> tuple[str, bytes | None]:
            try:
                response = await client.get(sm_url)
                response.raise_for_status()
                return ("ok", response.content or b"")
            except Exception:
                return ("fail", None)

        while sitemap_queue:
            if documents_fetched >= SITEMAP_MAX_DOCUMENT_FETCHES:
                truncated_by_doc_cap = True
                break

            batch: list[tuple[str, int]] = []
            while (
                len(batch) < SITEMAP_PARALLEL_FETCHES
                and sitemap_queue
                and documents_fetched + len(batch) < SITEMAP_MAX_DOCUMENT_FETCHES
            ):
                sm_url, depth = sitemap_queue.popleft()
                if sm_url in seen_sitemaps or depth > 8:
                    continue
                seen_sitemaps.add(sm_url)
                batch.append((sm_url, depth))

            if not batch:
                break

            documents_fetched += len(batch)
            fetched_tuples = await asyncio.gather(*[_fetch_sitemap_body(sm_url) for sm_url, _ in batch])

            hit_collect_cap = False
            for (_sm_url, depth), (fetch_st, content) in zip(batch, fetched_tuples, strict=True):
                if fetch_st != "ok" or content is None:
                    http_failure_count += 1
                    continue
                http_success_count += 1
                doc_parse_ok = False
                try:
                    for _event, elem in ET.iterparse(io.BytesIO(content), events=("end",)):
                        if len(all_page_urls) >= SITEMAP_MAX_PAGE_URLS_COLLECTED:
                            hit_collect_cap = True
                            truncated_by_url_collect_cap = True
                            break
                        local = _local_xml_tag(elem.tag)
                        if local == "url":
                            loc_text: str | None = None
                            for el in elem:
                                if _local_xml_tag(el.tag) == "loc" and el.text and el.text.strip():
                                    loc_text = el.text.strip()
                                    break
                            elem.clear()
                            if not loc_text:
                                continue
                            try:
                                norm = _normalize_url_string(loc_text)
                            except AppError:
                                continue
                            if norm not in seen_page_urls:
                                seen_page_urls.add(norm)
                                all_page_urls.append(norm)
                        elif local == "sitemap":
                            for el in elem:
                                if _local_xml_tag(el.tag) == "loc" and el.text and el.text.strip():
                                    raw_loc = el.text.strip()
                                    try:
                                        # Nested sitemap docs are never path-filtered; query must stay for Shopify.
                                        norm_loc = _normalize_sitemap_document_url(raw_loc)
                                    except AppError:
                                        continue
                                    sitemap_queue.append((norm_loc, depth + 1))
                            elem.clear()
                    doc_parse_ok = True
                except ET.ParseError:
                    xml_parse_failure_count += 1

                if doc_parse_ok:
                    had_successful_xml_document = True

                if hit_collect_cap:
                    break

            if progress:
                await progress(documents_fetched, len(all_page_urls))
            if hit_collect_cap:
                break

    urls_discovered_total = len(all_page_urls)
    filtered = _filter_discovered_page_urls(all_page_urls, include_rules, exclude_rules)
    urls_after_filter = len(filtered)
    hub_urls_added = 0
    if include_rules and filtered:
        filtered, hub_urls_added = await _expand_urls_from_matching_hubs(
            filtered,
            include_rules=include_rules,
            exclude_rules=exclude_rules,
            max_urls=max_urls,
        )
    final_urls = filtered[:max_urls]
    if include_rules and filtered:
        log.info(
            "sitemap_discovery.filtered",
            urls_discovered_total=urls_discovered_total,
            urls_after_filter=urls_after_filter,
            hub_urls_added=hub_urls_added,
            final_url_count=len(final_urls),
        )

    return SitemapDiscoveryResult(
        urls=final_urls,
        http_success_count=http_success_count,
        http_failure_count=http_failure_count,
        xml_parse_failure_count=xml_parse_failure_count,
        had_successful_xml_document=had_successful_xml_document,
        truncated_by_doc_cap=truncated_by_doc_cap or truncated_by_url_collect_cap,
        urls_discovered_total=urls_discovered_total,
        urls_after_filter=urls_after_filter,
        hub_urls_added=hub_urls_added,
    )


async def _collect_sitemap_urls(
    start_url: str,
    max_urls: int,
    include_rules: list[dict[str, str]],
    exclude_rules: list[dict[str, str]],
    *,
    progress: Callable[[int, int], Awaitable[None]] | None = None,
) -> list[str]:
    """Backward-compatible wrapper returning URL list only."""
    discovered = await _discover_sitemap_urls(
        start_url, max_urls, include_rules, exclude_rules, progress=progress
    )
    return discovered.urls


async def _fetch_pages_for_urls(
    urls: list[str],
    *,
    crawl_budget_bytes: int | None,
) -> tuple[list[dict[str, object]], int, str | None, int, str]:
    """Fetch each URL; optionally stop when extracted text bytes reach crawl budget."""
    pages: list[dict[str, object]] = []
    links_discovered = 0
    preview_image_url: str | None = None
    stored_text_bytes = 0
    stopped_reason = "complete"
    async with httpx.AsyncClient(timeout=20, follow_redirects=True, headers=_WEBSITE_CRAWL_HEADERS) as client:
        for url in urls:
            if crawl_budget_bytes is not None and stored_text_bytes >= crawl_budget_bytes and len(pages) > 0:
                stopped_reason = "budget"
                break
            try:
                response = await client.get(url)
                status_code = int(response.status_code)
                response.raise_for_status()
                _charged, html = _charge_bytes_and_html_from_response(
                    response, remaining_budget_bytes=_CRAWL_BODY_CHARGE_CAP_BYTES
                )
                if preview_image_url is None:
                    preview_image_url = _extract_social_preview_image(html, url)
                page_text = _extract_page_text(html)
                page_title = _extract_page_title(html)
                if crawl_budget_bytes is not None:
                    remaining_storage = max(0, crawl_budget_bytes - stored_text_bytes)
                    if remaining_storage <= 0 and len(pages) > 0:
                        stopped_reason = "budget"
                        break
                    page_text = _truncate_utf8_to_bytes(page_text, remaining_storage)
                text_bytes = len(page_text.encode("utf-8")) if page_text else 0
                stored_text_bytes += text_bytes
                pages.append(
                    {
                        "url": url,
                        "depth": 0,
                        "http_status": status_code,
                        "text": page_text,
                        "title": page_title,
                    }
                )
                links_discovered += len(_extract_links(url, html))
                if crawl_budget_bytes is not None and stored_text_bytes >= crawl_budget_bytes:
                    stopped_reason = "budget"
                    break
            except Exception:
                pages.append({"url": url, "depth": 0, "http_status": None, "text": "", "title": None})
    if stopped_reason == "complete":
        stopped_reason = "no_more_links"
    return pages, links_discovered, preview_image_url, stored_text_bytes, stopped_reason


def _website_page_index_metadata_json(page: dict[str, object]) -> str:
    """Persist extracted text on the page row so a later embedding-only worker pass can reload."""
    return json.dumps(
        {
            "extracted_text": str(page.get("text") or ""),
            "title": page.get("title"),
        },
        default=str,
    )


def _knowledge_source_page_row_params(page: dict[str, object], index: int) -> dict[str, str | int | None]:
    return {
        f"url_{index}": str(page["url"]),
        f"depth_{index}": int(page["depth"]),
        f"status_{index}": "parsed" if str(page["text"]).strip() else "failed",
        f"http_status_{index}": page["http_status"],
        f"metadata_{index}": _website_page_index_metadata_json(page),
    }


async def _bulk_upsert_knowledge_source_pages(
    db: AsyncSession,
    *,
    source_id: UUID,
    crawl_run_id: UUID,
    user_id: UUID,
    pages: list[dict[str, object]],
) -> None:
    if not pages:
        return
    for start in range(0, len(pages), DASHBOARD_PAGE_UPSERT_CHUNK):
        chunk = pages[start : start + DASHBOARD_PAGE_UPSERT_CHUNK]
        placeholders = ", ".join(
            f"(:source_id, :crawl_run_id, :user_id, :url_{i}, :depth_{i}, :status_{i}, :http_status_{i}, now(), cast(:metadata_{i} as jsonb))"
            for i in range(len(chunk))
        )
        params: dict[str, str | int | None] = {
            "source_id": str(source_id),
            "crawl_run_id": str(crawl_run_id),
            "user_id": str(user_id),
        }
        for i, page in enumerate(chunk):
            params.update(_knowledge_source_page_row_params(page, i))
        await db.execute(
            text(
                f"""
                insert into public.knowledge_source_pages (
                  knowledge_source_id, crawl_run_id, user_id, url, depth, status, http_status, last_crawled_at, metadata
                ) values {placeholders}
                on conflict (knowledge_source_id, url)
                do update set
                  crawl_run_id = excluded.crawl_run_id,
                  depth = excluded.depth,
                  status = excluded.status,
                  http_status = excluded.http_status,
                  last_crawled_at = now(),
                  metadata = excluded.metadata
                """
            ),
            params,
        )


async def _upsert_knowledge_source_page(
    db: AsyncSession,
    *,
    source_id: UUID,
    crawl_run_id: UUID,
    user_id: UUID,
    page: dict[str, object],
) -> None:
    await _bulk_upsert_knowledge_source_pages(
        db,
        source_id=source_id,
        crawl_run_id=crawl_run_id,
        user_id=user_id,
        pages=[page],
    )


async def _load_website_pages_for_embedding(
    db: AsyncSession, *, source_id: UUID, crawl_run_id: UUID, user_id: UUID
) -> list[dict[str, object]]:
    result = await db.execute(
        text(
            """
            select url, depth, http_status, status, metadata
            from public.knowledge_source_pages
            where knowledge_source_id = :sid
              and crawl_run_id = :rid
              and user_id = :uid
            order by url asc
            """
        ),
        {"sid": str(source_id), "rid": str(crawl_run_id), "uid": str(user_id)},
    )
    pages: list[dict[str, object]] = []
    for row in result.mappings().all():
        md_raw = row["metadata"]
        md = dict(md_raw or {}) if isinstance(md_raw, dict) else {}
        pages.append(
            {
                "url": row["url"],
                "depth": int(row["depth"] or 0),
                "http_status": row["http_status"],
                "text": str(md.get("extracted_text") or ""),
                "title": md.get("title"),
            }
        )
    return pages


async def _release_dashboard_job_for_embedding_pass(
    db: AsyncSession,
    *,
    job_id: UUID,
    embedding_handoff: dict[str, object],
) -> None:
    await db.execute(
        text(
            """
            update public.indexing_jobs
            set status = 'queued',
                phase = 'embedding_queued',
                progress_pct = greatest(progress_pct, 25),
                metrics = coalesce(metrics, '{}'::jsonb) || cast(:extra as jsonb)
            where id = :job_id
            """
        ),
        {"job_id": str(job_id), "extra": json.dumps(embedding_handoff, default=str)},
    )
    await db.commit()


async def _dashboard_update_crawl_job_progress(
    db: AsyncSession,
    *,
    job_id: UUID,
    pages_total_cap: int,
    pages_processed: int,
    crawl_http_bytes_so_far: int | None = None,
    crawl_phase: str = "fetching_html",
) -> None:
    """Update indexing job counters only (no per-page DB writes). Used while the worker fetches HTML."""
    pct = 5 + int(min(19, 19 * pages_processed / max(pages_total_cap, 1)))
    mextra: dict[str, object] = {"crawl_phase": crawl_phase}
    if crawl_http_bytes_so_far is not None:
        mextra["crawl_http_bytes"] = crawl_http_bytes_so_far
    await db.execute(
        text(
            """
            update public.indexing_jobs
            set pages_total = :pages_total,
                pages_processed = :pages_processed,
                progress_pct = :progress_pct,
                phase = 'crawling',
                metrics = coalesce(metrics, '{}'::jsonb) || cast(:mextra as jsonb)
            where id = :job_id
            """
        ),
        {
            "job_id": str(job_id),
            "pages_total": pages_total_cap,
            "pages_processed": pages_processed,
            "progress_pct": min(24, pct),
            "mextra": json.dumps(mextra),
        },
    )


async def _dashboard_persist_crawl_pages(
    db: AsyncSession,
    *,
    source_id: UUID,
    crawl_run_id: UUID,
    user_id: UUID,
    job_id: UUID,
    pages: list[dict[str, object]],
    pages_total_cap: int,
    crawl_http_bytes_so_far: int | None = None,
) -> None:
    """Bulk-upsert crawled page rows once per phase (chunked), then refresh job progress."""
    await _require_knowledge_source_exists(db, source_id)
    await _bulk_upsert_knowledge_source_pages(
        db,
        source_id=source_id,
        crawl_run_id=crawl_run_id,
        user_id=user_id,
        pages=pages,
    )
    await _dashboard_update_crawl_job_progress(
        db,
        job_id=job_id,
        pages_total_cap=pages_total_cap,
        pages_processed=len(pages),
        crawl_http_bytes_so_far=crawl_http_bytes_so_far,
        crawl_phase="fetching_html",
    )


async def _reset_website_source_pages_for_discovery(
    db: AsyncSession, *, source_id: UUID, user_id: UUID
) -> None:
    """Remove prior crawl rows so retrain / a new run does not keep an old 4-link plan."""
    await db.execute(
        text(
            """
            delete from public.knowledge_source_pages
            where knowledge_source_id = :source_id and user_id = :user_id
            """
        ),
        {"source_id": str(source_id), "user_id": str(user_id)},
    )


async def _dashboard_seed_queued_urls(
    db: AsyncSession,
    *,
    source_id: UUID,
    crawl_run_id: UUID,
    user_id: UUID,
    job_id: UUID,
    urls: list[str],
) -> None:
    """Insert placeholder rows so the dashboard shows link count before HTML is fetched."""
    sid = str(source_id)
    rid = str(crawl_run_id)
    uid = str(user_id)
    jid = str(job_id)
    extra_base = {"crawl_phase": "urls_discovered", "discovery": "sitemap"}
    if not urls:
        await db.execute(
            text(
                """
                update public.indexing_jobs
                set pages_total = 0, pages_processed = 0, progress_pct = 4, phase = 'crawling',
                    metrics = coalesce(metrics, '{}'::jsonb) || cast(:extra as jsonb)
                where id = :job_id
                """
            ),
            {"job_id": jid, "extra": json.dumps({**extra_base, "seeding_total": 0})},
        )
        return

    n = len(urls)
    chunk_idx = 0
    for start in range(0, n, DASHBOARD_SEED_URL_INSERT_CHUNK):
        chunk = urls[start : start + DASHBOARD_SEED_URL_INSERT_CHUNK]
        placeholders = ", ".join(
            f"(:source_id, :crawl_run_id, :user_id, :url_{i}, 0, 'queued', null, now())"
            for i in range(len(chunk))
        )
        params: dict[str, str] = {
            "source_id": sid,
            "crawl_run_id": rid,
            "user_id": uid,
        }
        for i, u in enumerate(chunk):
            params[f"url_{i}"] = u
        await db.execute(
            text(
                f"""
                insert into public.knowledge_source_pages (
                  knowledge_source_id, crawl_run_id, user_id, url, depth, status, http_status, last_crawled_at
                ) values {placeholders}
                on conflict (knowledge_source_id, url)
                do update set
                  crawl_run_id = excluded.crawl_run_id,
                  depth = excluded.depth,
                  status = 'queued',
                  http_status = null,
                  last_crawled_at = now()
                """
            ),
            params,
        )
        if chunk_idx % 3 == 0:
            done = min(start + len(chunk), n)
            await db.execute(
                text(
                    """
                    update public.indexing_jobs
                    set metrics = coalesce(metrics, '{}'::jsonb) || cast(:extra as jsonb)
                    where id = :job_id
                    """
                ),
                {
                    "job_id": jid,
                    "extra": json.dumps(
                        {
                            **extra_base,
                            "seeding_done": done,
                            "seeding_total": n,
                            "crawl_phase": "seeding_urls",
                        }
                    ),
                },
            )
        chunk_idx += 1

    await db.execute(
        text(
            """
            update public.indexing_jobs
            set pages_total = :pages_total, pages_processed = 0, progress_pct = 4, phase = 'crawling',
                metrics = coalesce(metrics, '{}'::jsonb) || cast(:extra as jsonb)
            where id = :job_id
            """
        ),
        {
            "job_id": jid,
            "pages_total": n,
            "extra": json.dumps({**extra_base, "seeding_done": n, "seeding_total": n}),
        },
    )


async def _exclude_remaining_queued_pages_for_run(
    db: AsyncSession, *, knowledge_source_id: UUID, crawl_run_id: UUID
) -> None:
    """Sitemap seeding inserts many `queued` placeholders; drop unfetched URLs only when the crawl fully finished."""
    await db.execute(
        text(
            """
            update public.knowledge_source_pages
            set status = 'excluded'::public.crawl_page_status
            where knowledge_source_id = :sid
              and crawl_run_id = :rid
              and status = 'queued'::public.crawl_page_status
            """
        ),
        {"sid": str(knowledge_source_id), "rid": str(crawl_run_id)},
    )


def _effective_crawl_budget_bytes(*, origin: str, remaining_storage: int, crawl_cap: int) -> int:
    """Onboarding uses the full knowledge pool so the first crawl can cover the whole sitemap (within plan storage)."""
    if origin == "onboarding":
        return max(0, remaining_storage)
    return min(max(0, remaining_storage), max(0, crawl_cap))


async def _count_queued_pages_for_source(db: AsyncSession, source_id: UUID, user_id: UUID) -> int:
    row = (
        await db.execute(
            text(
                """
                select count(*)::int as c
                from public.knowledge_source_pages
                where knowledge_source_id = :sid and user_id = :uid and status = 'queued'::public.crawl_page_status
                """
            ),
            {"sid": str(source_id), "uid": str(user_id)},
        )
    ).mappings().one()
    return int(row["c"] or 0)


async def _list_queued_page_urls(
    db: AsyncSession, source_id: UUID, user_id: UUID, *, limit: int
) -> list[str]:
    rows = (
        await db.execute(
            text(
                """
                select url
                from public.knowledge_source_pages
                where knowledge_source_id = :sid and user_id = :uid and status = 'queued'::public.crawl_page_status
                order by url asc
                limit :limit
                """
            ),
            {"sid": str(source_id), "uid": str(user_id), "limit": max(1, int(limit))},
        )
    ).mappings().all()
    return [str(r["url"]) for r in rows]


async def _source_has_pending_indexing_job(db: AsyncSession, source_id: UUID, user_id: UUID) -> bool:
    row = (
        await db.execute(
            text(
                """
                select id
                from public.indexing_jobs
                where knowledge_source_id = :sid and user_id = :uid and status in ('queued', 'running')
                limit 1
                """
            ),
            {"sid": str(source_id), "uid": str(user_id)},
        )
    ).first()
    return row is not None


async def _clear_crawl_continuation_metadata(db: AsyncSession, source_id: UUID, user_id: UUID) -> None:
    await db.execute(
        text(
            """
            update public.knowledge_sources
            set metadata = coalesce(metadata, '{}'::jsonb) - 'fetch_queued_only' - 'continuation_crawl_run_id'
            where id = :sid and user_id = :uid
            """
        ),
        {"sid": str(source_id), "uid": str(user_id)},
    )


async def _maybe_enqueue_crawl_continuation(
    db: AsyncSession,
    *,
    user_id: UUID,
    source: KnowledgeSourceDTO,
    crawl_run_id: UUID,
    crawl_stopped_reason: str,
) -> None:
    if crawl_stopped_reason not in ("budget", "max_pages_safety"):
        await _clear_crawl_continuation_metadata(db, source.id, user_id)
        return
    queued = await _count_queued_pages_for_source(db, source.id, user_id)
    if queued <= 0:
        await _clear_crawl_continuation_metadata(db, source.id, user_id)
        return
    if await _source_has_pending_indexing_job(db, source.id, user_id):
        return
    md = dict(source.metadata or {})
    md["fetch_queued_only"] = True
    md["continuation_crawl_run_id"] = str(crawl_run_id)
    await db.execute(
        text(
            """
            update public.knowledge_sources
            set metadata = cast(:metadata as jsonb)
            where id = :sid and user_id = :uid
            """
        ),
        {"sid": str(source.id), "uid": str(user_id), "metadata": json.dumps(md)},
    )
    await enqueue_index_website_source_queued(db, source.id, user_id)
    await db.commit()


async def _dashboard_fetch_planned_urls_in_batches(
    db: AsyncSession,
    *,
    job_id: UUID,
    source: KnowledgeSourceDTO,
    user_id: UUID,
    crawl_run_id: UUID,
    urls: list[str],
    crawl_budget_bytes: int,
) -> tuple[list[dict[str, object]], int, str | None, int, str]:
    """Fetch HTML for URLs discovered via sitemap in small batches; commit after each batch."""
    if not urls:
        return [], 0, None, 0, "complete"
    total_saved = 0
    all_pages: list[dict[str, object]] = []
    preview: str | None = None
    stopped = "complete"
    links_discovered = 0
    n = len(urls)
    for start in range(0, n, DASHBOARD_CRAWL_CONTENT_BATCH):
        remaining_budget = max(0, crawl_budget_bytes - total_saved)
        if remaining_budget == 0 and len(all_pages) > 0:
            stopped = "budget"
            break
        batch_urls = urls[start : start + DASHBOARD_CRAWL_CONTENT_BATCH]
        batch_pages, ld, pv, b_used, st = await _fetch_pages_for_urls(
            batch_urls, crawl_budget_bytes=remaining_budget
        )
        total_saved += b_used
        links_discovered += ld
        if preview is None and pv:
            preview = pv
        if st == "budget":
            stopped = "budget"
        for p in batch_pages:
            p.setdefault("depth", 0)
        all_pages.extend(batch_pages)
        await _dashboard_update_crawl_job_progress(
            db,
            job_id=job_id,
            pages_total_cap=n,
            pages_processed=len(all_pages),
            crawl_http_bytes_so_far=total_saved,
        )
        await db.commit()
        if stopped == "budget":
            break
    if all_pages:
        await _dashboard_persist_crawl_pages(
            db,
            source_id=source.id,
            crawl_run_id=crawl_run_id,
            user_id=user_id,
            job_id=job_id,
            pages=all_pages,
            pages_total_cap=n,
            crawl_http_bytes_so_far=total_saved,
        )
        await db.commit()
    if stopped not in ("budget", "max_pages_safety"):
        await _exclude_remaining_queued_pages_for_run(
            db, knowledge_source_id=source.id, crawl_run_id=crawl_run_id
        )
        await db.commit()
    return all_pages, links_discovered, preview, total_saved, stopped


def _indexing_failure_row(exc: BaseException) -> tuple[str, dict[str, object]]:
    """Short `error_message` text and structured `metrics.failure` for debugging."""
    if isinstance(exc, AppError):
        short = f"{exc.code}: {exc.message}"[:1000]
        detail: dict[str, object] = {
            "kind": "app_error",
            "code": exc.code,
            "message": exc.message,
            "details": exc.details or {},
        }
    else:
        short = f"{type(exc).__name__}: {exc}"[:1000]
        tb = "".join(traceback.format_exception(exc))
        if len(tb) > 12_000:
            tb = f"{tb[:12_000]}\n... (traceback truncated)"
        detail = {"kind": "unexpected", "exc_type": type(exc).__name__, "message": str(exc), "traceback": tb}
    return short, {"failure": detail}


async def _finalize_indexing_failure(
    db: AsyncSession,
    *,
    job_id: UUID,
    source_id: UUID,
    crawl_run_id: UUID | None,
    exc: BaseException,
) -> None:
    short, metrics_obj = _indexing_failure_row(exc)
    metrics_json = json.dumps(metrics_obj, default=str)
    await db.rollback()
    if crawl_run_id is not None:
        await _exclude_remaining_queued_pages_for_run(
            db, knowledge_source_id=source_id, crawl_run_id=crawl_run_id
        )
    await db.execute(
        text(
            """
            update public.knowledge_sources
            set status = 'failed', error_message = :error_message
            where id = :source_id
            """
        ),
        {"source_id": str(source_id), "error_message": short[:1000]},
    )
    await db.execute(
        text(
            """
            update public.indexing_jobs
            set status = 'failed', phase = 'failed', finished_at = now(), error_message = :error_message,
                metrics = cast(:metrics as jsonb)
            where id = :job_id
            """
        ),
        {"job_id": str(job_id), "error_message": short[:1000], "metrics": metrics_json},
    )
    if crawl_run_id is not None:
        await db.execute(
            text(
                """
                update public.knowledge_crawl_runs
                set status = 'failed', error_message = :error_message, finished_at = now()
                where id = :crawl_run_id
                """
            ),
                {"crawl_run_id": str(crawl_run_id), "error_message": short[:1000]},
            )
    src_meta = (
        await db.execute(
            text(
                """
                select user_id, agent_id, type, title
                from public.knowledge_sources
                where id = cast(:sid as uuid)
                """
            ),
            {"sid": str(source_id)},
        )
    ).mappings().first()
    await db.commit()
    if src_meta:
        from app.domains.notifications.links import href_knowledge_source
        from app.domains.notifications.service import create_notification_best_effort

        uid = src_meta["user_id"]
        aid = src_meta["agent_id"]
        stype = str(src_meta["type"])
        stitle = str(src_meta["title"] or "Knowledge source")
        await create_notification_best_effort(
            db,
            user_id=uid,
            kind="knowledge_index_failed",
            title="Knowledge indexing failed",
            body=f"Could not finish indexing “{stitle}”: {short[:280]}",
            href=href_knowledge_source(source_type=stype, agent_id=aid, source_id=source_id),
            metadata={"source_id": str(source_id), "source_type": stype},
        )


async def record_worker_indexing_surrogate_failure(
    db: AsyncSession, job_id: UUID, user_id: UUID, exc: BaseException
) -> None:
    """If the worker caught an exception after the job was left `running`, persist a failure row."""
    short, metrics_obj = _indexing_failure_row(exc)
    worker_note = f"(worker) {short}"[:1000]
    metrics_obj["failure"]["worker_caught"] = True
    metrics_json = json.dumps(metrics_obj, default=str)
    await db.execute(
        text(
            """
            update public.indexing_jobs j
            set status = 'failed', phase = 'failed', finished_at = coalesce(j.finished_at, now()),
                error_message = coalesce(j.error_message, :error_message),
                metrics = coalesce(j.metrics, '{}'::jsonb) || cast(:metrics as jsonb)
            where j.id = :job_id and j.user_id = :user_id and j.status = 'running'
            """
        ),
        {"job_id": str(job_id), "user_id": str(user_id), "error_message": worker_note, "metrics": metrics_json},
    )
    await db.execute(
        text(
            """
            update public.knowledge_sources s
            set status = 'failed', error_message = coalesce(s.error_message, :error_message)
            from public.indexing_jobs j
            where s.id = j.knowledge_source_id
              and j.id = :job_id and j.user_id = :user_id
              and j.status = 'failed'
              and s.status = 'indexing'
            """
        ),
        {"job_id": str(job_id), "user_id": str(user_id), "error_message": worker_note},
    )
    await db.execute(
        text(
            """
            update public.knowledge_crawl_runs r
            set status = 'failed', error_message = coalesce(r.error_message, :error_message),
                finished_at = coalesce(r.finished_at, now())
            where r.id = (
              select cr.id
              from public.knowledge_crawl_runs cr
              where cr.knowledge_source_id = (
                select knowledge_source_id from public.indexing_jobs where id = :job_id and user_id = :user_id
              )
                and cr.status = 'running'
              order by cr.started_at desc nulls last
              limit 1
            )
            """
        ),
        {"job_id": str(job_id), "user_id": str(user_id), "error_message": worker_note},
    )
    await db.commit()


async def _complete_website_indexing_embedding_phase(
    db: AsyncSession,
    *,
    job_id: UUID,
    user_id: UUID,
    source: KnowledgeSourceDTO,
    crawl_run_id: UUID,
    preview_image_url: str | None,
    usable_pages: list[dict[str, object]],
    urls_fetched: int,
    crawl_http_bytes: int,
    crawl_budget_bytes: int,
    crawl_stopped_reason: str,
    website_discovery_mode: str,
    links_discovered: int,
    fetch_stats: dict[str, object],
    partial_warnings: list[dict[str, object]],
    dashboard_planned_url_count: int | None,
    include_rules: list[dict[str, str]],
    exclude_rules: list[dict[str, str]],
    effective_storage_cap_bytes: int,
) -> str | None:
    crawled_pages_with_text_count = len(usable_pages)
    skipped_already_indexed = 0

    md = dict(source.metadata or {})
    origin = str(md.get("origin") or "")
    website_mode = str(md.get("website_mode") or "crawl")
    should_dedupe_already_indexed = origin == "dashboard_website" and website_mode != "individual"

    if should_dedupe_already_indexed and usable_pages:
        page_key_pairs: list[tuple[dict[str, object], str]] = []
        url_keys: list[str] = []
        for page in usable_pages:
            url = str(page.get("url") or "").strip()
            if not url:
                continue
            key = _url_duplicate_key(url)
            page_key_pairs.append((page, key))
            url_keys.append(key)

        url_keys = list({k for k in url_keys if k})
        if url_keys:
            existing_keys_rows = (
                await db.execute(
                    text(
                        """
                        select distinct
                          regexp_replace(lower(btrim(c.metadata->>'page_url')), '/+$', '') as url_key
                        from public.knowledge_chunks c
                        where c.agent_id = cast(:agent_id as uuid)
                          and c.knowledge_source_id <> cast(:source_id as uuid)
                          and c.metadata ? 'page_url'
                          and regexp_replace(lower(btrim(c.metadata->>'page_url')), '/+$', '') = any(
                            cast(:url_keys as text[])
                          )
                        """
                    ),
                    {
                        "agent_id": str(source.agent_id),
                        "source_id": str(source.id),
                        "url_keys": url_keys,
                    },
                )
            ).mappings().all()
            existing_keys = {str(r.get("url_key") or "") for r in existing_keys_rows if r.get("url_key") is not None}

            before = len(usable_pages)
            usable_pages = [p for (p, k) in page_key_pairs if k not in existing_keys]
            skipped_already_indexed = before - len(usable_pages)

    indexed_source_bytes = sum(len(str(p["text"]).encode("utf-8")) for p in usable_pages)
    chunk_records: list[dict[str, object]] = []
    for page in usable_pages:
        page_url = str(page.get("url") or source.source_url or "")
        page_text = str(page.get("text") or "")
        page_title = str(page.get("title") or source.title or "").strip()
        prefix_parts = [p for p in [page_title, page_url] if p]
        chunk_prefix = " | ".join(prefix_parts).strip()
        for page_chunk_index, chunk in enumerate(_chunk_text(page_text)):
            chunk_with_context = f"{chunk_prefix}\n\n{chunk}".strip() if chunk_prefix else chunk
            chunk_records.append(
                {
                    "content": chunk_with_context,
                    "raw_content": chunk,
                    "page_url": page_url,
                    "page_title": page_title,
                    "page_chunk_index": page_chunk_index,
                }
            )

    if not chunk_records:
        # If all candidate pages were already indexed elsewhere for this agent,
        # embedding would produce zero chunks. Mark the source/job succeeded anyway.
        planned_urls = dashboard_planned_url_count if dashboard_planned_url_count is not None else urls_fetched

        storage_stopped_reason = "already_indexed" if skipped_already_indexed > 0 else "complete"
        success_metrics_obj: dict[str, object] = {
            "chunk_count": 0,
            "indexed_source_bytes": 0,
            "crawl_http_bytes": crawl_http_bytes,
            "crawl_budget_bytes": crawl_budget_bytes,
            "crawl_stopped_reason": crawl_stopped_reason,
            "storage_stopped_reason": storage_stopped_reason,
            "urls_planned": planned_urls,
            "urls_fetched": urls_fetched,
            "urls_indexed": 0,
            "urls_skipped_already_indexed": skipped_already_indexed,
            "discovery_mode": website_discovery_mode,
            "filter_summary": _website_filter_summary(include_rules, exclude_rules),
            "links_discovered_raw_anchors": links_discovered,
            "fetch_stats": fetch_stats,
        }
        if partial_warnings:
            success_metrics_obj["warnings"] = partial_warnings
        success_metrics = json.dumps(success_metrics_obj, default=str)

        await db.execute(
            text(
                """
                update public.knowledge_sources
                set status = 'ready', last_indexed_at = now(), error_message = null
                where id = :source_id
                """
            ),
            {"source_id": str(source.id)},
        )
        await db.execute(
            text(
                """
                update public.indexing_jobs
                set status = 'succeeded', phase = 'complete', progress_pct = 100, finished_at = now(),
                    pages_total = :pages_total,
                    pages_processed = :pages_processed,
                    chunks_embedded = 0,
                    metrics = cast(:metrics as jsonb)
                where id = :job_id
                """
            ),
            {
                "job_id": str(job_id),
                "pages_total": planned_urls,
                "pages_processed": urls_fetched,
                "metrics": success_metrics,
            },
        )
        await db.execute(
            text(
                """
                update public.knowledge_crawl_runs
                set status = 'succeeded',
                    pages_discovered = :pages_discovered,
                    pages_crawled = :pages_crawled,
                    pages_failed = :pages_failed,
                    links_discovered = :links_discovered,
                    finished_at = now()
                where id = :crawl_run_id
                """
            ),
            {
                "crawl_run_id": str(crawl_run_id),
                "pages_discovered": urls_fetched,
                "pages_crawled": crawled_pages_with_text_count,
                "pages_failed": urls_fetched - crawled_pages_with_text_count,
                "links_discovered": links_discovered,
            },
        )
        await db.commit()

        from app.domains.notifications.links import href_knowledge_source
        from app.domains.notifications.service import create_notification_best_effort

        await create_notification_best_effort(
            db,
            user_id=user_id,
            kind="knowledge_index_complete",
            title="Knowledge indexing complete",
            body=f"“{source.title}” is indexed and ready for answers.",
            href=href_knowledge_source(
                source_type=source.type, agent_id=source.agent_id, source_id=source.id
            ),
            metadata={"source_id": str(source.id), "source_type": source.type},
        )

        return preview_image_url

    best_title = next((str(p.get("title") or "").strip() for p in usable_pages if str(p.get("title") or "").strip()), "")
    if best_title:
        await db.execute(
            text(
                """
                update public.knowledge_sources
                set title = :title
                where id = :source_id
                """
            ),
            {"source_id": str(source.id), "title": best_title[:255]},
        )

    await db.execute(
        text(
            """
            update public.indexing_jobs
            set phase = 'chunking', chunks_total = :chunks_total, progress_pct = 45
            where id = :job_id
            """
        ),
        {"job_id": str(job_id), "chunks_total": len(chunk_records)},
    )

    embeddings = await _embed_texts([str(rec["content"]) for rec in chunk_records])
    await db.execute(
        text(
            """
            update public.indexing_jobs
            set phase = 'embedding', chunks_embedded = :chunks_embedded, progress_pct = 75
            where id = :job_id
            """
        ),
        {"job_id": str(job_id), "chunks_embedded": len(embeddings)},
    )
    await db.execute(
        text("delete from public.knowledge_chunks where knowledge_source_id = :source_id"),
        {"source_id": str(source.id)},
    )

    chunks_persisted = 0
    storage_stopped_reason = "complete"
    for batch_start in range(0, len(chunk_records), _CHUNK_PERSIST_BATCH_SIZE):
        batch_end = min(batch_start + _CHUNK_PERSIST_BATCH_SIZE, len(chunk_records))
        for idx in range(batch_start, batch_end):
            rec = chunk_records[idx]
            chunk = str(rec["content"])
            embedding = embeddings[idx]
            if len(embedding) != EMBEDDING_DIMENSION:
                raise AppError(
                    code="knowledge.embedding_dimension_mismatch",
                    message="Indexed content format does not match storage",
                    status_code=500,
                    details={"expected": EMBEDDING_DIMENSION, "actual": len(embedding)},
                )
            await db.execute(
                text(
                    """
                    insert into public.knowledge_chunks (
                      agent_id, user_id, knowledge_source_id, chunk_index, content, embedding, token_count, metadata
                    ) values (
                      :agent_id, :user_id, :knowledge_source_id, :chunk_index, :content, CAST(:embedding AS vector), :token_count, CAST(:metadata AS jsonb)
                    )
                    """
                ),
                {
                    "agent_id": str(source.agent_id),
                    "user_id": str(user_id),
                    "knowledge_source_id": str(source.id),
                    "chunk_index": idx,
                    "content": chunk,
                    "embedding": _vector_literal(embedding),
                    "token_count": _token_estimate(chunk),
                    "metadata": json.dumps(
                        {
                            "source_url": source.source_url,
                            "page_url": rec.get("page_url"),
                            "page_title": rec.get("page_title"),
                            "raw_content": rec.get("raw_content"),
                            "chunk_index": idx,
                            "page_chunk_index": rec.get("page_chunk_index"),
                        }
                    ),
                },
            )
            chunks_persisted += 1
        await db.commit()
        used_after_batch = await _agent_used_storage_bytes(db, user_id=user_id, agent_id=source.agent_id)
        if used_after_batch >= effective_storage_cap_bytes:
            storage_stopped_reason = "saved_storage_budget"
            break

    await db.execute(
        text(
            """
            update public.knowledge_sources
            set status = 'ready', last_indexed_at = now(), error_message = null
            where id = :source_id
            """
        ),
        {"source_id": str(source.id)},
    )
    planned_urls = dashboard_planned_url_count if dashboard_planned_url_count is not None else urls_fetched
    success_metrics_obj: dict[str, object] = {
        "chunk_count": chunks_persisted,
        "indexed_source_bytes": indexed_source_bytes,
        "crawl_http_bytes": crawl_http_bytes,
        "crawl_budget_bytes": crawl_budget_bytes,
        "crawl_stopped_reason": crawl_stopped_reason,
        "storage_stopped_reason": storage_stopped_reason,
        "urls_planned": planned_urls,
        "urls_fetched": urls_fetched,
        "urls_indexed": len(usable_pages),
        "discovery_mode": website_discovery_mode,
        "filter_summary": _website_filter_summary(include_rules, exclude_rules),
        "links_discovered_raw_anchors": links_discovered,
        "fetch_stats": fetch_stats,
    }
    if skipped_already_indexed:
        success_metrics_obj["urls_skipped_already_indexed"] = skipped_already_indexed
    if partial_warnings:
        success_metrics_obj["warnings"] = partial_warnings
    success_metrics = json.dumps(success_metrics_obj, default=str)
    await db.execute(
        text(
            """
            update public.indexing_jobs
            set status = 'succeeded', phase = 'complete', progress_pct = 100, finished_at = now(),
                pages_total = :pages_total,
                pages_processed = :pages_processed,
                chunks_embedded = :chunks_embedded,
                metrics = cast(:metrics as jsonb)
            where id = :job_id
            """
        ),
        {
            "job_id": str(job_id),
            "pages_total": planned_urls,
            "pages_processed": urls_fetched,
            "chunks_embedded": chunks_persisted,
            "metrics": success_metrics,
        },
    )
    await db.execute(
        text(
            """
            update public.knowledge_crawl_runs
            set status = 'succeeded',
                pages_discovered = :pages_discovered,
                pages_crawled = :pages_crawled,
                pages_failed = :pages_failed,
                links_discovered = :links_discovered,
                finished_at = now()
            where id = :crawl_run_id
            """
        ),
        {
            "crawl_run_id": str(crawl_run_id),
            "pages_discovered": urls_fetched,
            "pages_crawled": crawled_pages_with_text_count,
            "pages_failed": urls_fetched - crawled_pages_with_text_count,
            "links_discovered": links_discovered,
        },
    )
    await db.commit()
    from app.domains.notifications.links import href_knowledge_source
    from app.domains.notifications.service import create_notification_best_effort

    await _maybe_enqueue_crawl_continuation(
        db,
        user_id=user_id,
        source=source,
        crawl_run_id=crawl_run_id,
        crawl_stopped_reason=crawl_stopped_reason,
    )
    queued_after = await _count_queued_pages_for_source(db, source.id, user_id)
    if queued_after <= 0:
        await create_notification_best_effort(
            db,
            user_id=user_id,
            kind="knowledge_index_complete",
            title="Knowledge indexing complete",
            body=f"“{source.title}” is indexed and ready for answers.",
            href=href_knowledge_source(
                source_type=source.type, agent_id=source.agent_id, source_id=source.id
            ),
            metadata={"source_id": str(source.id), "source_type": source.type},
        )
    return preview_image_url


async def _process_website_embedding_only(
    db: AsyncSession,
    job_id: UUID,
    user_id: UUID,
    source: KnowledgeSourceDTO,
    job_metrics: dict[str, object],
) -> str | None:
    raw_cr = job_metrics.get("embedding_crawl_run_id")
    if raw_cr is None:
        raise AppError(
            code="knowledge.embedding_handoff_missing",
            message="Indexing job is missing crawl handoff metadata; try re-indexing the source.",
            status_code=500,
        )
    crawl_run_id = UUID(str(raw_cr))
    md = dict(source.metadata or {})
    _mode, _max_pages, include_rules, exclude_rules = _ingest_params_from_metadata(md)

    pages = await _load_website_pages_for_embedding(
        db, source_id=source.id, crawl_run_id=crawl_run_id, user_id=user_id
    )
    usable_pages = [p for p in pages if str(p["text"]).strip()]
    if not usable_pages:
        fs = job_metrics.get("fetch_stats")
        fetch_stats = dict(fs) if isinstance(fs, dict) else {}
        planned_urls_metrics = (
            int(job_metrics["dashboard_planned_url_count"])
            if job_metrics.get("dashboard_planned_url_count") is not None
            else len(pages)
        )
        raise AppError(
            code="knowledge.scrape_empty",
            message="No usable text from fetched pages (all empty, failed HTTP, or blocked by crawl budget).",
            status_code=422,
            details={
                "seed_url": source.source_url,
                "knowledge_source_id": str(source.id),
                "discovery_mode": str(job_metrics.get("website_discovery_mode") or "unknown"),
                "urls_planned": planned_urls_metrics,
                **fetch_stats,
                "crawl_stopped_reason": str(job_metrics.get("crawl_stopped_reason") or ""),
                "filter_summary": _website_filter_summary(include_rules, exclude_rules),
            },
        )

    _plan_slug, _, plan_features = await _fetch_active_subscription_plan(db, user_id)
    included_storage_cap_bytes = _included_storage_bytes_from_plan_features(plan_features)
    effective_storage_cap_bytes = _effective_storage_cap_bytes(included_storage_cap_bytes)

    pv = job_metrics.get("preview_image_url")
    preview_image_url = str(pv) if pv not in (None, "") else None

    crawl_http_bytes = int(job_metrics.get("crawl_http_bytes") or 0)
    crawl_budget_bytes = int(job_metrics.get("crawl_budget_bytes") or 0)
    crawl_stopped_reason = str(job_metrics.get("crawl_stopped_reason") or "complete")
    website_discovery_mode = str(job_metrics.get("website_discovery_mode") or "unknown")
    links_discovered = int(job_metrics.get("links_discovered") or 0)
    fs2 = job_metrics.get("fetch_stats")
    fetch_stats = dict(fs2) if isinstance(fs2, dict) else {}
    pw = job_metrics.get("partial_warnings")
    partial_warnings: list[dict[str, object]] = (
        [dict(x) for x in pw if isinstance(x, dict)] if isinstance(pw, list) else []
    )
    dpc = job_metrics.get("dashboard_planned_url_count")
    dashboard_planned_url_count = int(dpc) if dpc is not None else None
    urls_fetched = int(job_metrics.get("urls_fetched") or len(pages))

    return await _complete_website_indexing_embedding_phase(
        db,
        job_id=job_id,
        user_id=user_id,
        source=source,
        crawl_run_id=crawl_run_id,
        preview_image_url=preview_image_url,
        usable_pages=usable_pages,
        urls_fetched=urls_fetched,
        crawl_http_bytes=crawl_http_bytes,
        crawl_budget_bytes=crawl_budget_bytes,
        crawl_stopped_reason=crawl_stopped_reason,
        website_discovery_mode=website_discovery_mode,
        links_discovered=links_discovered,
        fetch_stats=fetch_stats,
        partial_warnings=partial_warnings,
        dashboard_planned_url_count=dashboard_planned_url_count,
        include_rules=include_rules,
        exclude_rules=exclude_rules,
        effective_storage_cap_bytes=effective_storage_cap_bytes,
    )


async def process_indexing_job(db: AsyncSession, job_id: UUID, user_id: UUID) -> str | None:
    result = await db.execute(
        text(
            """
            select
              j.id as job_id,
              j.phase::text as job_phase,
              j.metrics as job_metrics,
              s.id, s.agent_id, s.user_id, s.type, s.title, s.status, s.source_url, s.storage_bucket, s.storage_path,
              s.metadata, s.error_message, s.last_indexed_at, s.created_at, s.updated_at
            from public.indexing_jobs j
            join public.knowledge_sources s on s.id = j.knowledge_source_id
            where j.id = :job_id and j.user_id = :user_id
            """
        ),
        {"job_id": str(job_id), "user_id": str(user_id)},
    )
    row = result.mappings().first()
    if row is None:
        raise AppError(code="knowledge.job_not_found", message="Indexing job not found", status_code=404)
    row_map = dict(row)
    row_map.pop("job_id", None)
    job_phase = str(row_map.pop("job_phase"))
    raw_metrics = row_map.pop("job_metrics")
    job_metrics_snapshot: dict[str, object] = dict(raw_metrics) if isinstance(raw_metrics, dict) else {}
    source = KnowledgeSourceDTO.model_validate(row_map)
    if source.type != "website" or not source.source_url:
        raise AppError(code="validation.invalid_input", message="Only website sources are supported", status_code=422)

    if job_phase == "embedding":
        cr_emb: UUID | None = None
        try:
            raw_handoff = job_metrics_snapshot.get("embedding_crawl_run_id")
            if raw_handoff is not None:
                cr_emb = UUID(str(raw_handoff))
        except ValueError:
            cr_emb = None
        try:
            return await _process_website_embedding_only(db, job_id, user_id, source, job_metrics_snapshot)
        except Exception as exc:
            if isinstance(exc, IntegrityError):
                detail = str(getattr(exc, "orig", None) or exc)
                if "knowledge_source_pages_knowledge_source_id_fkey" in detail:
                    exc = AppError(
                        code="knowledge.source_removed",
                        message="This website source was deleted while indexing was still running.",
                        status_code=409,
                    )
            await _finalize_indexing_failure(
                db,
                job_id=job_id,
                source_id=source.id,
                crawl_run_id=cr_emb,
                exc=exc,
            )
            if isinstance(exc, AppError):
                raise
            raise AppError(code="knowledge.indexing_failed", message="Indexing job failed", status_code=500) from exc

    md = dict(source.metadata or {})
    mode, max_pages, include_rules, exclude_rules = _ingest_params_from_metadata(md)

    crawl_run_id: UUID | None = None
    preview_image_url: str | None = None

    try:
        crawl_settings: dict[str, object] = {
            "max_pages": max_pages,
            "website_mode": mode,
            "include_rules": include_rules,
            "exclude_rules": exclude_rules,
        }
        _plan_slug, _, plan_features = await _fetch_active_subscription_plan(db, user_id)
        included_storage_cap_bytes = _included_storage_bytes_from_plan_features(plan_features)
        effective_storage_cap_bytes = _effective_storage_cap_bytes(included_storage_cap_bytes)
        currently_used_storage_bytes = await _agent_used_storage_bytes(
            db, user_id=user_id, agent_id=source.agent_id
        )
        remaining_storage = max(0, effective_storage_cap_bytes - currently_used_storage_bytes)
        crawl_cap = website_crawl_cap_bytes(plan_features, included_storage_cap_bytes)
        origin = str(md.get("origin") or "")
        crawl_budget_bytes = _effective_crawl_budget_bytes(
            origin=origin,
            remaining_storage=remaining_storage,
            crawl_cap=crawl_cap,
        )
        pages_persisted_incrementally = False
        dashboard_planned_url_count: int | None = None
        website_discovery_mode = "unknown"
        fetch_queued_only = bool(md.get("fetch_queued_only")) and md.get("continuation_crawl_run_id")

        if fetch_queued_only:
            crawl_run_id = UUID(str(md["continuation_crawl_run_id"]))
            queued_urls = await _list_queued_page_urls(db, source.id, user_id, limit=max_pages)
            website_discovery_mode = "resume_queued"
            if queued_urls:
                dashboard_planned_url_count = await _count_queued_pages_for_source(db, source.id, user_id)
                pages, links_discovered, preview_image_url, crawl_http_bytes, crawl_stopped_reason = (
                    await _dashboard_fetch_planned_urls_in_batches(
                        db,
                        job_id=job_id,
                        source=source,
                        user_id=user_id,
                        crawl_run_id=crawl_run_id,
                        urls=queued_urls,
                        crawl_budget_bytes=crawl_budget_bytes,
                    )
                )
                pages_persisted_incrementally = True
            else:
                await _clear_crawl_continuation_metadata(db, source.id, user_id)
                pages = []
                links_discovered = 0
                crawl_http_bytes = 0
                crawl_stopped_reason = "complete"
        else:
            crawl_run_id = await _create_crawl_run(db, source, user_id, crawl_settings)
            await _reset_website_source_pages_for_discovery(db, source_id=source.id, user_id=user_id)
            await db.commit()

        async def _indexing_sitemap_progress(docs_scanned: int, urls_in_sitemap: int) -> None:
            pct = min(
                20,
                5
                + int(
                    15
                    * min(docs_scanned, SITEMAP_MAX_DOCUMENT_FETCHES)
                    / max(SITEMAP_MAX_DOCUMENT_FETCHES, 1)
                ),
            )
            await db.execute(
                text(
                    """
                    update public.indexing_jobs
                    set progress_pct = greatest(progress_pct, :pct),
                        metrics = coalesce(metrics, '{}'::jsonb) || cast(:extra as jsonb)
                    where id = :job_id
                    """
                ),
                {
                    "job_id": str(job_id),
                    "pct": pct,
                    "extra": json.dumps(
                        {
                            "crawl_phase": "sitemap_discovery",
                            "sitemap_docs_fetched": docs_scanned,
                            "sitemap_urls_discovered": urls_in_sitemap,
                        }
                    ),
                },
            )
            await db.commit()

        async def _indexing_apply_filters_progress(discovered: SitemapDiscoveryResult) -> None:
            await db.execute(
                text(
                    """
                    update public.indexing_jobs
                    set progress_pct = greatest(progress_pct, 22),
                        metrics = coalesce(metrics, '{}'::jsonb) || cast(:extra as jsonb)
                    where id = :job_id
                    """
                ),
                {
                    "job_id": str(job_id),
                    "extra": json.dumps(
                        {
                            "crawl_phase": "applying_path_filters",
                            "sitemap_urls_discovered": discovered.urls_discovered_total,
                            "sitemap_urls_after_filter": discovered.urls_after_filter,
                            "sitemap_hub_urls_added": discovered.hub_urls_added,
                            "sitemap_urls_matched": len(discovered.urls),
                        }
                    ),
                },
            )
            await db.commit()

        if fetch_queued_only:
            pass
        elif mode == "individual":
            website_discovery_mode = "individual"
            pages, links_discovered, preview_image_url, crawl_http_bytes, crawl_stopped_reason = await _fetch_pages_for_urls(
                [source.source_url], crawl_budget_bytes=crawl_budget_bytes
            )
        elif mode == "sitemap":
            website_discovery_mode = "sitemap"
            sitemap_discovered = await _discover_sitemap_urls(
                source.source_url,
                max_pages,
                include_rules,
                exclude_rules,
                progress=_indexing_sitemap_progress,
            )
            await _indexing_apply_filters_progress(sitemap_discovered)
            sitemap_urls = _cap_path_filtered_urls(
                sitemap_discovered.urls, include_rules, exclude_rules, max_pages
            )
            if not sitemap_urls:
                raise _app_error_for_empty_sitemap_discovery(sitemap_discovered)
            dashboard_planned_url_count = len(sitemap_urls)
            if _website_ingest_uses_live_discovery(md):
                await _dashboard_seed_queued_urls(
                    db,
                    source_id=source.id,
                    crawl_run_id=crawl_run_id,
                    user_id=user_id,
                    job_id=job_id,
                    urls=sitemap_urls,
                )
                await db.commit()
                pages, links_discovered, preview_image_url, crawl_http_bytes, crawl_stopped_reason = (
                    await _dashboard_fetch_planned_urls_in_batches(
                        db,
                        job_id=job_id,
                        source=source,
                        user_id=user_id,
                        crawl_run_id=crawl_run_id,
                        urls=sitemap_urls,
                        crawl_budget_bytes=crawl_budget_bytes,
                    )
                )
                pages_persisted_incrementally = True
            else:
                pages, links_discovered, preview_image_url, crawl_http_bytes, crawl_stopped_reason = await _fetch_pages_for_urls(
                    sitemap_urls, crawl_budget_bytes=crawl_budget_bytes
                )
        else:
            if _website_ingest_uses_live_discovery(md):
                sitemap_discovered = await _discover_sitemap_urls(
                    source.source_url,
                    max_pages,
                    include_rules,
                    exclude_rules,
                    progress=_indexing_sitemap_progress,
                )
                await _indexing_apply_filters_progress(sitemap_discovered)
                sitemap_plan = _cap_path_filtered_urls(
                    sitemap_discovered.urls, include_rules, exclude_rules, max_pages
                )
                if len(sitemap_plan) > 0:
                    website_discovery_mode = "sitemap"
                    dashboard_planned_url_count = len(sitemap_plan)
                    await _dashboard_seed_queued_urls(
                        db,
                        source_id=source.id,
                        crawl_run_id=crawl_run_id,
                        user_id=user_id,
                        job_id=job_id,
                        urls=sitemap_plan,
                    )
                    await db.commit()
                    pages, links_discovered, preview_image_url, crawl_http_bytes, crawl_stopped_reason = (
                        await _dashboard_fetch_planned_urls_in_batches(
                            db,
                            job_id=job_id,
                            source=source,
                            user_id=user_id,
                            crawl_run_id=crawl_run_id,
                            urls=sitemap_plan,
                            crawl_budget_bytes=crawl_budget_bytes,
                        )
                    )
                    pages_persisted_incrementally = True
                else:
                    website_discovery_mode = "bfs"
                    await db.execute(
                        text(
                            """
                            update public.indexing_jobs
                            set pages_total = 0, pages_processed = 0, progress_pct = 4, phase = 'crawling',
                                metrics = coalesce(metrics, '{}'::jsonb) || cast(:extra as jsonb)
                            where id = :job_id
                            """
                        ),
                        {
                            "job_id": str(job_id),
                            "extra": json.dumps({"crawl_phase": "bfs_fetch", "discovery": "bfs"}),
                        },
                    )
                    await db.commit()

                    async def _bfs_progress(
                        snapshot: list[dict[str, object]],
                        _bytes_so_far: int,
                        crawl_visits: int = 0,
                    ) -> None:
                        if snapshot:
                            await _bulk_upsert_knowledge_source_pages(
                                db,
                                source_id=source.id,
                                crawl_run_id=crawl_run_id,
                                user_id=user_id,
                                pages=snapshot,
                            )
                        cumulative = max(len(snapshot), crawl_visits)
                        await _dashboard_update_crawl_job_progress(
                            db,
                            job_id=job_id,
                            pages_total_cap=max_pages,
                            pages_processed=cumulative,
                            crawl_http_bytes_so_far=_bytes_so_far,
                            crawl_phase="bfs_fetch",
                        )
                        await db.commit()

                    bfs_progress_every = (
                        1 if md.get("origin") == "onboarding" else DASHBOARD_CRAWL_PROGRESS_EVERY
                    )
                    pages, links_discovered, preview_image_url, crawl_http_bytes, crawl_stopped_reason, bfs_visits = (
                        await _crawl_pages(
                            source.source_url,
                            max_pages,
                            include_rules,
                            exclude_rules,
                            crawl_budget_bytes=crawl_budget_bytes,
                            progress_hook=_bfs_progress,
                            progress_every=bfs_progress_every,
                        )
                    )
                    await _bfs_progress(pages, crawl_http_bytes, bfs_visits)
                    if pages:
                        await _dashboard_persist_crawl_pages(
                            db,
                            source_id=source.id,
                            crawl_run_id=crawl_run_id,
                            user_id=user_id,
                            job_id=job_id,
                            pages=pages,
                            pages_total_cap=max_pages,
                            crawl_http_bytes_so_far=crawl_http_bytes,
                        )
                        await db.commit()
                    pages_persisted_incrementally = True

        if fetch_queued_only and not pages:
            await _clear_crawl_continuation_metadata(db, source.id, user_id)
            await db.execute(
                text(
                    """
                    update public.indexing_jobs
                    set status = 'succeeded', phase = 'complete', progress_pct = 100, finished_at = now()
                    where id = :job_id
                    """
                ),
                {"job_id": str(job_id)},
            )
            await db.execute(
                text(
                    """
                    update public.knowledge_sources
                    set status = 'ready', last_indexed_at = now(), error_message = null
                    where id = :source_id
                    """
                ),
                {"source_id": str(source.id)},
            )
            await db.commit()
            return preview_image_url

        usable_pages = [p for p in pages if str(p["text"]).strip()]
        fetch_stats = _website_fetch_page_stats(pages)
        planned_urls_metrics = (
            dashboard_planned_url_count if dashboard_planned_url_count is not None else len(pages)
        )
        if not usable_pages:
            raise AppError(
                code="knowledge.scrape_empty",
                message="No usable text from fetched pages (all empty, failed HTTP, or blocked by crawl budget).",
                status_code=422,
                details={
                    "seed_url": source.source_url,
                    "knowledge_source_id": str(source.id),
                    "discovery_mode": website_discovery_mode,
                    "urls_planned": planned_urls_metrics,
                    **fetch_stats,
                    "crawl_stopped_reason": crawl_stopped_reason,
                    "filter_summary": _website_filter_summary(include_rules, exclude_rules),
                },
            )

        partial_warnings: list[dict[str, object]] = []
        if len(usable_pages) < len(pages):
            partial_warnings.append(
                {
                    "code": "partial_empty_text",
                    "indexed_urls": len(usable_pages),
                    "skipped_empty_urls": len(pages) - len(usable_pages),
                }
            )

        split_after_crawl = md.get("origin") in ("dashboard_website", "onboarding")

        if pages_persisted_incrementally:
            await db.execute(
                text(
                    """
                    update public.indexing_jobs
                    set progress_pct = 25,
                        metrics = coalesce(metrics, '{}'::jsonb) || cast(:extra as jsonb)
                    where id = :job_id
                    """
                ),
                {
                    "job_id": str(job_id),
                    "extra": json.dumps(
                        {
                            "crawl_http_bytes": crawl_http_bytes,
                            "crawl_stopped_reason": crawl_stopped_reason,
                            "urls_fetched": len(pages),
                            "discovery_mode": website_discovery_mode,
                            **fetch_stats,
                        }
                    ),
                },
            )
        else:
            await db.execute(
                text(
                    """
                    update public.indexing_jobs
                    set pages_total = :pages_total, pages_processed = :pages_processed, progress_pct = 25
                    where id = :job_id
                    """
                ),
                {"job_id": str(job_id), "pages_total": len(pages), "pages_processed": len(pages)},
            )

        if not pages_persisted_incrementally:
            await _dashboard_persist_crawl_pages(
                db,
                source_id=source.id,
                crawl_run_id=crawl_run_id,
                user_id=user_id,
                job_id=job_id,
                pages=pages,
                pages_total_cap=len(pages),
                crawl_http_bytes_so_far=crawl_http_bytes,
            )

        await db.commit()

        if split_after_crawl:
            embedding_handoff: dict[str, object] = {
                "embedding_crawl_run_id": str(crawl_run_id),
                "preview_image_url": preview_image_url,
                "dashboard_planned_url_count": dashboard_planned_url_count,
                "website_discovery_mode": website_discovery_mode,
                "crawl_http_bytes": crawl_http_bytes,
                "crawl_budget_bytes": crawl_budget_bytes,
                "crawl_stopped_reason": crawl_stopped_reason,
                "links_discovered": links_discovered,
                "fetch_stats": fetch_stats,
                "partial_warnings": partial_warnings,
                "urls_fetched": len(pages),
            }
            await _release_dashboard_job_for_embedding_pass(db, job_id=job_id, embedding_handoff=embedding_handoff)
            return preview_image_url

        return await _complete_website_indexing_embedding_phase(
            db,
            job_id=job_id,
            user_id=user_id,
            source=source,
            crawl_run_id=crawl_run_id,
            preview_image_url=preview_image_url,
            usable_pages=usable_pages,
            urls_fetched=len(pages),
            crawl_http_bytes=crawl_http_bytes,
            crawl_budget_bytes=crawl_budget_bytes,
            crawl_stopped_reason=crawl_stopped_reason,
            website_discovery_mode=website_discovery_mode,
            links_discovered=links_discovered,
            fetch_stats=fetch_stats,
            partial_warnings=partial_warnings,
            dashboard_planned_url_count=dashboard_planned_url_count,
            include_rules=include_rules,
            exclude_rules=exclude_rules,
            effective_storage_cap_bytes=effective_storage_cap_bytes,
        )
    except Exception as exc:
        if isinstance(exc, IntegrityError):
            detail = str(getattr(exc, "orig", None) or exc)
            if "knowledge_source_pages_knowledge_source_id_fkey" in detail:
                exc = AppError(
                    code="knowledge.source_removed",
                    message="This website source was deleted while indexing was still running.",
                    status_code=409,
                )
        await _finalize_indexing_failure(
            db,
            job_id=job_id,
            source_id=source.id,
            crawl_run_id=crawl_run_id,
            exc=exc,
        )
        if isinstance(exc, AppError):
            raise
        raise AppError(code="knowledge.indexing_failed", message="Indexing job failed", status_code=500) from exc


async def index_website_source(
    db: AsyncSession, source_id: UUID, user_id: UUID
) -> tuple[KnowledgeSourceDTO, IndexJobDTO]:
    source, job = await enqueue_index_website_source(db, source_id, user_id, mark_running=True)
    await process_indexing_job(db, job.id, user_id)
    refreshed = await get_latest_job(db, source_id, user_id)
    source_out = await _load_source(db, source_id, user_id)
    return source_out, refreshed or job


async def index_file_source(
    db: AsyncSession, source_id: UUID, user_id: UUID
) -> tuple[KnowledgeSourceDTO, IndexJobDTO]:
    source = await _load_source(db, source_id, user_id)
    if source.type != "file":
        raise AppError(code="validation.invalid_input", message="Only file sources are supported", status_code=422)
    if not source.storage_bucket or not source.storage_path:
        raise AppError(code="validation.invalid_input", message="File source is missing storage reference", status_code=422)

    job = await _create_job(db, source, user_id)
    await _set_job_running(db, job.id)
    await db.execute(
        text(
            """
            update public.knowledge_sources
            set status = 'indexing', error_message = null
            where id = :source_id
            """
        ),
        {"source_id": str(source.id)},
    )
    await db.commit()

    try:
        metadata = dict(source.metadata or {})
        text_content = str(metadata.get("extracted_text") or metadata.get("text") or metadata.get("content") or "").strip()
        if not text_content:
            raise AppError(
                code="knowledge.file_text_missing",
                message="File text is not available for indexing. Re-upload with extracted text metadata.",
                status_code=422,
            )

        _, _, plan_features = await _fetch_active_subscription_plan(db, user_id)
        included_storage_cap_bytes = _included_storage_bytes_from_plan_features(plan_features)
        effective_storage_cap_bytes = _effective_storage_cap_bytes(included_storage_cap_bytes)
        currently_used_storage_bytes = await _agent_used_storage_bytes(
            db, user_id=user_id, agent_id=source.agent_id
        )
        remaining_budget_bytes = max(0, effective_storage_cap_bytes - currently_used_storage_bytes)
        indexed_text = _truncate_utf8_to_bytes(text_content, remaining_budget_bytes)
        if not indexed_text:
            raise AppError(
                code="knowledge.storage_budget_exhausted",
                message="Knowledge storage limit reached. Upgrade to add more file content.",
                status_code=422,
            )

        chunks = _chunk_text(indexed_text)
        if not chunks:
            raise AppError(code="knowledge.chunking_empty", message="No chunks were produced from file text", status_code=422)

        await db.execute(
            text(
                """
                update public.indexing_jobs
                set phase = 'chunking', chunks_total = :chunks_total, progress_pct = 45
                where id = :job_id
                """
            ),
            {"job_id": str(job.id), "chunks_total": len(chunks)},
        )

        embeddings = await _embed_texts(chunks)
        await db.execute(
            text(
                """
                update public.indexing_jobs
                set phase = 'embedding', chunks_embedded = :chunks_embedded, progress_pct = 75
                where id = :job_id
                """
            ),
            {"job_id": str(job.id), "chunks_embedded": len(embeddings)},
        )
        await db.execute(
            text("delete from public.knowledge_chunks where knowledge_source_id = :source_id"),
            {"source_id": str(source.id)},
        )

        for idx, chunk in enumerate(chunks):
            embedding = embeddings[idx]
            if len(embedding) != EMBEDDING_DIMENSION:
                raise AppError(
                    code="knowledge.embedding_dimension_mismatch",
                    message="Indexed content format does not match storage",
                    status_code=500,
                    details={"expected": EMBEDDING_DIMENSION, "actual": len(embedding)},
                )
            await db.execute(
                text(
                    """
                    insert into public.knowledge_chunks (
                      agent_id, user_id, knowledge_source_id, chunk_index, content, embedding, token_count, metadata
                    ) values (
                      :agent_id, :user_id, :knowledge_source_id, :chunk_index, :content, CAST(:embedding AS vector), :token_count, CAST(:metadata AS jsonb)
                    )
                    """
                ),
                {
                    "agent_id": str(source.agent_id),
                    "user_id": str(user_id),
                    "knowledge_source_id": str(source.id),
                    "chunk_index": idx,
                    "content": chunk,
                    "embedding": _vector_literal(embedding),
                    "token_count": _token_estimate(chunk),
                    "metadata": json.dumps(
                        {
                            "source_type": "file",
                            "storage_bucket": source.storage_bucket,
                            "storage_path": source.storage_path,
                            "chunk_index": idx,
                        }
                    ),
                },
            )

        await db.execute(
            text(
                """
                update public.knowledge_sources
                set status = 'ready', last_indexed_at = now(), error_message = null
                where id = :source_id
                """
            ),
            {"source_id": str(source.id)},
        )
        await db.execute(
            text(
                """
                update public.indexing_jobs
                set status = 'succeeded', phase = 'complete', progress_pct = 100, finished_at = now(),
                    chunks_embedded = :chunks_embedded,
                    metrics = cast(:metrics as jsonb)
                where id = :job_id
                """
            ),
            {
                "job_id": str(job.id),
                "chunks_embedded": len(chunks),
                "metrics": json.dumps(
                    {
                        "chunk_count": len(chunks),
                        "indexed_source_bytes": len(indexed_text.encode("utf-8")),
                        "storage_stopped_reason": "complete",
                    }
                ),
            },
        )
        await db.commit()
        from app.domains.notifications.links import href_knowledge_source
        from app.domains.notifications.service import create_notification_best_effort

        await create_notification_best_effort(
            db,
            user_id=user_id,
            kind="knowledge_index_complete",
            title="Knowledge indexing complete",
            body=f"“{source.title}” is indexed and ready for answers.",
            href=href_knowledge_source(
                source_type=source.type, agent_id=source.agent_id, source_id=source.id
            ),
            metadata={"source_id": str(source.id), "source_type": source.type},
        )
    except Exception as exc:
        await _finalize_indexing_failure(
            db,
            job_id=job.id,
            source_id=source.id,
            crawl_run_id=None,
            exc=exc,
        )
        if isinstance(exc, AppError):
            raise
        raise AppError(code="knowledge.indexing_failed", message="Indexing job failed", status_code=500) from exc

    refreshed = await get_latest_job(db, source_id, user_id)
    source_out = await _load_source(db, source_id, user_id)
    return source_out, refreshed or job


async def index_text_snippet_source(
    db: AsyncSession, source_id: UUID, user_id: UUID
) -> tuple[KnowledgeSourceDTO, IndexJobDTO]:
    source = await _load_source(db, source_id, user_id)
    if source.type != "text_snippet":
        raise AppError(code="validation.invalid_input", message="Only text_snippet sources are supported", status_code=422)

    job = await _create_job(db, source, user_id)
    await _set_job_running(db, job.id)
    await db.execute(
        text(
            """
            update public.knowledge_sources
            set status = 'indexing', error_message = null
            where id = :source_id
            """
        ),
        {"source_id": str(source.id)},
    )
    await db.commit()

    try:
        text_content = str(source.raw_text or "").strip()
        if not text_content:
            raise AppError(
                code="knowledge.snippet_text_missing",
                message="Snippet text is empty; add body text before indexing.",
                status_code=422,
            )

        _, _, plan_features = await _fetch_active_subscription_plan(db, user_id)
        included_storage_cap_bytes = _included_storage_bytes_from_plan_features(plan_features)
        effective_storage_cap_bytes = _effective_storage_cap_bytes(included_storage_cap_bytes)
        currently_used_storage_bytes = await _agent_used_storage_bytes(
            db, user_id=user_id, agent_id=source.agent_id
        )
        remaining_budget_bytes = max(0, effective_storage_cap_bytes - currently_used_storage_bytes)
        indexed_text = _truncate_utf8_to_bytes(text_content, remaining_budget_bytes)
        if not indexed_text:
            raise AppError(
                code="knowledge.storage_budget_exhausted",
                message="Knowledge storage limit reached. Upgrade to add more snippet content.",
                status_code=422,
            )

        chunks = _chunk_text(indexed_text)
        if not chunks:
            raise AppError(code="knowledge.chunking_empty", message="No chunks were produced from snippet text", status_code=422)

        await db.execute(
            text(
                """
                update public.indexing_jobs
                set phase = 'chunking', chunks_total = :chunks_total, progress_pct = 45
                where id = :job_id
                """
            ),
            {"job_id": str(job.id), "chunks_total": len(chunks)},
        )

        embeddings = await _embed_texts(chunks)
        await db.execute(
            text(
                """
                update public.indexing_jobs
                set phase = 'embedding', chunks_embedded = :chunks_embedded, progress_pct = 75
                where id = :job_id
                """
            ),
            {"job_id": str(job.id), "chunks_embedded": len(embeddings)},
        )
        await db.execute(
            text("delete from public.knowledge_chunks where knowledge_source_id = :source_id"),
            {"source_id": str(source.id)},
        )

        for idx, chunk in enumerate(chunks):
            embedding = embeddings[idx]
            if len(embedding) != EMBEDDING_DIMENSION:
                raise AppError(
                    code="knowledge.embedding_dimension_mismatch",
                    message="Indexed content format does not match storage",
                    status_code=500,
                    details={"expected": EMBEDDING_DIMENSION, "actual": len(embedding)},
                )
            await db.execute(
                text(
                    """
                    insert into public.knowledge_chunks (
                      agent_id, user_id, knowledge_source_id, chunk_index, content, embedding, token_count, metadata
                    ) values (
                      :agent_id, :user_id, :knowledge_source_id, :chunk_index, :content, CAST(:embedding AS vector), :token_count, CAST(:metadata AS jsonb)
                    )
                    """
                ),
                {
                    "agent_id": str(source.agent_id),
                    "user_id": str(user_id),
                    "knowledge_source_id": str(source.id),
                    "chunk_index": idx,
                    "content": chunk,
                    "embedding": _vector_literal(embedding),
                    "token_count": _token_estimate(chunk),
                    "metadata": json.dumps(
                        {
                            "source_type": "text_snippet",
                            "snippet_title": source.title,
                            "chunk_index": idx,
                        }
                    ),
                },
            )

        await db.execute(
            text(
                """
                update public.knowledge_sources
                set status = 'ready', last_indexed_at = now(), error_message = null
                where id = :source_id
                """
            ),
            {"source_id": str(source.id)},
        )
        await db.execute(
            text(
                """
                update public.indexing_jobs
                set status = 'succeeded', phase = 'complete', progress_pct = 100, finished_at = now(),
                    chunks_embedded = :chunks_embedded,
                    metrics = cast(:metrics as jsonb)
                where id = :job_id
                """
            ),
            {
                "job_id": str(job.id),
                "chunks_embedded": len(chunks),
                "metrics": json.dumps(
                    {
                        "chunk_count": len(chunks),
                        "indexed_source_bytes": len(indexed_text.encode("utf-8")),
                        "storage_stopped_reason": "complete",
                    }
                ),
            },
        )
        await db.commit()
        from app.domains.notifications.links import href_knowledge_source
        from app.domains.notifications.service import create_notification_best_effort

        await create_notification_best_effort(
            db,
            user_id=user_id,
            kind="knowledge_index_complete",
            title="Knowledge indexing complete",
            body=f"“{source.title}” is indexed and ready for answers.",
            href=href_knowledge_source(
                source_type=source.type, agent_id=source.agent_id, source_id=source.id
            ),
            metadata={"source_id": str(source.id), "source_type": source.type},
        )
    except Exception as exc:
        await _finalize_indexing_failure(
            db,
            job_id=job.id,
            source_id=source.id,
            crawl_run_id=None,
            exc=exc,
        )
        if isinstance(exc, AppError):
            raise
        raise AppError(code="knowledge.indexing_failed", message="Indexing job failed", status_code=500) from exc

    refreshed = await get_latest_job(db, source_id, user_id)
    source_out = await _load_source(db, source_id, user_id)
    return source_out, refreshed or job


def _qa_pair_index_text(question: str, answer: str) -> str:
    q = (question or "").strip()
    a = (answer or "").strip()
    return f"Question: {q}\n\nAnswer: {a}"


async def index_qa_source(db: AsyncSession, source_id: UUID, user_id: UUID) -> tuple[KnowledgeSourceDTO, IndexJobDTO]:
    source = await _load_source(db, source_id, user_id)
    if source.type != "q_and_a":
        raise AppError(code="validation.invalid_input", message="Only q_and_a sources are supported", status_code=422)

    qa_row = (
        await db.execute(
            text(
                """
                select q.question, q.answer
                from public.knowledge_qa_items q
                join public.knowledge_sources s on s.id = q.knowledge_source_id
                where q.knowledge_source_id = :source_id and s.user_id = :user_id and s.type = 'q_and_a'
                order by q.created_at asc
                limit 1
                """
            ),
            {"source_id": str(source.id), "user_id": str(user_id)},
        )
    ).mappings().first()
    if qa_row is None:
        raise AppError(
            code="knowledge.qa_row_missing",
            message="Q&A pair has no question/answer row to index.",
            status_code=422,
        )

    job = await _create_job(db, source, user_id)
    await _set_job_running(db, job.id)
    await db.execute(
        text(
            """
            update public.knowledge_sources
            set status = 'indexing', error_message = null
            where id = :source_id
            """
        ),
        {"source_id": str(source.id)},
    )
    await db.commit()

    try:
        text_content = _normalize_text(_qa_pair_index_text(str(qa_row["question"]), str(qa_row["answer"])))
        if not text_content:
            raise AppError(
                code="knowledge.qa_text_missing",
                message="Q&A text is empty; add a question and answer before indexing.",
                status_code=422,
            )

        _, _, plan_features = await _fetch_active_subscription_plan(db, user_id)
        included_storage_cap_bytes = _included_storage_bytes_from_plan_features(plan_features)
        effective_storage_cap_bytes = _effective_storage_cap_bytes(included_storage_cap_bytes)
        currently_used_storage_bytes = await _agent_used_storage_bytes(
            db, user_id=user_id, agent_id=source.agent_id
        )
        remaining_budget_bytes = max(0, effective_storage_cap_bytes - currently_used_storage_bytes)
        indexed_text = _truncate_utf8_to_bytes(text_content, remaining_budget_bytes)
        if not indexed_text:
            raise AppError(
                code="knowledge.storage_budget_exhausted",
                message="Knowledge storage limit reached. Upgrade to add more Q&A content.",
                status_code=422,
            )

        chunks = _chunk_text(indexed_text)
        if not chunks:
            raise AppError(code="knowledge.chunking_empty", message="No chunks were produced from Q&A text", status_code=422)

        await db.execute(
            text(
                """
                update public.indexing_jobs
                set phase = 'chunking', chunks_total = :chunks_total, progress_pct = 45
                where id = :job_id
                """
            ),
            {"job_id": str(job.id), "chunks_total": len(chunks)},
        )

        embeddings = await _embed_texts(chunks)
        await db.execute(
            text(
                """
                update public.indexing_jobs
                set phase = 'embedding', chunks_embedded = :chunks_embedded, progress_pct = 75
                where id = :job_id
                """
            ),
            {"job_id": str(job.id), "chunks_embedded": len(embeddings)},
        )
        await db.execute(
            text("delete from public.knowledge_chunks where knowledge_source_id = :source_id"),
            {"source_id": str(source.id)},
        )

        for idx, chunk in enumerate(chunks):
            embedding = embeddings[idx]
            if len(embedding) != EMBEDDING_DIMENSION:
                raise AppError(
                    code="knowledge.embedding_dimension_mismatch",
                    message="Indexed content format does not match storage",
                    status_code=500,
                    details={"expected": EMBEDDING_DIMENSION, "actual": len(embedding)},
                )
            await db.execute(
                text(
                    """
                    insert into public.knowledge_chunks (
                      agent_id, user_id, knowledge_source_id, chunk_index, content, embedding, token_count, metadata
                    ) values (
                      :agent_id, :user_id, :knowledge_source_id, :chunk_index, :content, CAST(:embedding AS vector), :token_count, CAST(:metadata AS jsonb)
                    )
                    """
                ),
                {
                    "agent_id": str(source.agent_id),
                    "user_id": str(user_id),
                    "knowledge_source_id": str(source.id),
                    "chunk_index": idx,
                    "content": chunk,
                    "embedding": _vector_literal(embedding),
                    "token_count": _token_estimate(chunk),
                    "metadata": json.dumps(
                        {
                            "source_type": "q_and_a",
                            "chunk_index": idx,
                        }
                    ),
                },
            )

        await db.execute(
            text(
                """
                update public.knowledge_sources
                set status = 'ready', last_indexed_at = now(), error_message = null
                where id = :source_id
                """
            ),
            {"source_id": str(source.id)},
        )
        await db.execute(
            text(
                """
                update public.indexing_jobs
                set status = 'succeeded', phase = 'complete', progress_pct = 100, finished_at = now(),
                    chunks_embedded = :chunks_embedded,
                    metrics = cast(:metrics as jsonb)
                where id = :job_id
                """
            ),
            {
                "job_id": str(job.id),
                "chunks_embedded": len(chunks),
                "metrics": json.dumps(
                    {
                        "chunk_count": len(chunks),
                        "indexed_source_bytes": len(indexed_text.encode("utf-8")),
                        "storage_stopped_reason": "complete",
                    }
                ),
            },
        )
        await db.commit()
        from app.domains.notifications.links import href_knowledge_source
        from app.domains.notifications.service import create_notification_best_effort

        await create_notification_best_effort(
            db,
            user_id=user_id,
            kind="knowledge_index_complete",
            title="Knowledge indexing complete",
            body=f"“{source.title}” is indexed and ready for answers.",
            href=href_knowledge_source(
                source_type=source.type, agent_id=source.agent_id, source_id=source.id
            ),
            metadata={"source_id": str(source.id), "source_type": source.type},
        )
    except Exception as exc:
        await _finalize_indexing_failure(
            db,
            job_id=job.id,
            source_id=source.id,
            crawl_run_id=None,
            exc=exc,
        )
        if isinstance(exc, AppError):
            raise
        raise AppError(code="knowledge.indexing_failed", message="Indexing job failed", status_code=500) from exc

    refreshed = await get_latest_job(db, source_id, user_id)
    source_out = await _load_source(db, source_id, user_id)
    return source_out, refreshed or job


async def create_and_index_text_snippet(
    db: AsyncSession, *, user_id: UUID, agent_id: UUID, title: str, snippet_text: str
) -> tuple[KnowledgeSourceDTO, IndexJobDTO]:
    body = (snippet_text or "").strip()
    if not body:
        raise AppError(code="validation.invalid_input", message="Snippet text cannot be empty", status_code=422)
    t = (title or "").strip()[:255]
    if not t:
        raise AppError(code="validation.invalid_input", message="Snippet title cannot be empty", status_code=422)

    _, _, plan_features = await _fetch_active_subscription_plan(db, user_id)
    included_storage_cap_bytes = _included_storage_bytes_from_plan_features(plan_features)
    effective_storage_cap_bytes = _effective_storage_cap_bytes(included_storage_cap_bytes)
    currently_used_storage_bytes = await _agent_used_storage_bytes(db, user_id=user_id, agent_id=agent_id)
    remaining_budget_bytes = max(0, effective_storage_cap_bytes - currently_used_storage_bytes)
    if remaining_budget_bytes <= 0:
        raise AppError(
            code="knowledge.storage_budget_exhausted",
            message="Knowledge storage limit reached. Remove or delete indexed content, or upgrade your plan.",
            status_code=422,
        )

    source = await create_source(
        db,
        user_id,
        KnowledgeSourceCreateRequest(
            agent_id=agent_id,
            type="text_snippet",
            title=t,
            raw_text=body,
            metadata={"origin": "dashboard_text_snippet"},
        ),
    )
    return await index_text_snippet_source(db, source.id, user_id)


async def create_and_index_qa_pair(
    db: AsyncSession, *, user_id: UUID, agent_id: UUID, question: str, answer: str
) -> tuple[KnowledgeSourceDTO, IndexJobDTO]:
    q = (question or "").strip()
    a = (answer or "").strip()
    if not q or not a:
        raise AppError(code="validation.invalid_input", message="Question and answer are required", status_code=422)
    title = q[:255] if q else "Q&A"

    _, _, plan_features = await _fetch_active_subscription_plan(db, user_id)
    included_storage_cap_bytes = _included_storage_bytes_from_plan_features(plan_features)
    effective_storage_cap_bytes = _effective_storage_cap_bytes(included_storage_cap_bytes)
    currently_used_storage_bytes = await _agent_used_storage_bytes(db, user_id=user_id, agent_id=agent_id)
    remaining_budget_bytes = max(0, effective_storage_cap_bytes - currently_used_storage_bytes)
    if remaining_budget_bytes <= 0:
        raise AppError(
            code="knowledge.storage_budget_exhausted",
            message="Knowledge storage limit reached. Remove or delete indexed content, or upgrade your plan.",
            status_code=422,
        )

    source = await create_source(
        db,
        user_id,
        KnowledgeSourceCreateRequest(
            agent_id=agent_id,
            type="q_and_a",
            title=title,
            metadata={"origin": "dashboard_qa"},
        ),
    )
    try:
        await db.execute(
            text(
                """
                insert into public.knowledge_qa_items (knowledge_source_id, question, answer)
                values (:knowledge_source_id, :question, :answer)
                """
            ),
            {"knowledge_source_id": str(source.id), "question": q, "answer": a},
        )
        await db.commit()
    except Exception:
        await db.execute(
            text("delete from public.knowledge_sources where id = :id and user_id = :user_id"),
            {"id": str(source.id), "user_id": str(user_id)},
        )
        await db.commit()
        raise

    return await index_qa_source(db, source.id, user_id)


async def create_failed_uploaded_file_source(
    db: AsyncSession,
    *,
    user_id: UUID,
    agent_id: UUID,
    filename: str,
    content_type: str | None,
    uploaded_bytes: int,
    error_message: str,
) -> FileUploadResultDTO:
    source = await create_source(
        db,
        user_id,
        KnowledgeSourceCreateRequest(
            agent_id=agent_id,
            type="file",
            title=(filename or "Uploaded file").strip()[:255] or "Uploaded file",
            storage_bucket="inline-upload",
            storage_path=_safe_storage_path(filename),
            status="failed",
            metadata={
                "origin": "dashboard_file_upload",
                "filename": filename,
                "content_type": content_type,
                "uploaded_bytes": uploaded_bytes,
            },
        ),
    )
    await db.execute(
        text(
            """
            update public.knowledge_sources
            set error_message = :error_message, updated_at = now()
            where id = :source_id and user_id = :user_id
            """
        ),
        {"source_id": str(source.id), "user_id": str(user_id), "error_message": error_message},
    )
    await db.commit()
    failed = await _load_source(db, source.id, user_id)
    return FileUploadResultDTO(source=failed, job=None, status="failed", error_message=failed.error_message)


async def create_and_index_uploaded_file(
    db: AsyncSession,
    *,
    user_id: UUID,
    agent_id: UUID,
    filename: str,
    content_type: str | None,
    payload: bytes,
) -> FileUploadResultDTO:
    if not payload:
        raise AppError(code="knowledge.file_empty", message="Uploaded file is empty", status_code=422)
    if len(payload) > 50 * 1024 * 1024:
        raise AppError(code="knowledge.file_too_large", message="Max file size is 50MB", status_code=422)

    extracted = _normalize_text(_extract_text_from_file_bytes(filename, content_type, payload))
    if not extracted:
        raise AppError(code="knowledge.file_text_missing", message="Could not extract readable text from file", status_code=422)

    _, _, plan_features = await _fetch_active_subscription_plan(db, user_id)
    included_storage_cap_bytes = _included_storage_bytes_from_plan_features(plan_features)
    effective_storage_cap_bytes = _effective_storage_cap_bytes(included_storage_cap_bytes)
    currently_used_storage_bytes = await _agent_used_storage_bytes(db, user_id=user_id, agent_id=agent_id)
    remaining_budget_bytes = max(0, effective_storage_cap_bytes - currently_used_storage_bytes)
    if remaining_budget_bytes <= 0:
        return await create_failed_uploaded_file_source(
            db,
            user_id=user_id,
            agent_id=agent_id,
            filename=filename,
            content_type=content_type,
            uploaded_bytes=len(payload),
            error_message="Knowledge storage limit reached. Remove indexed content or upgrade your plan.",
        )

    extracted_bytes = len(extracted.encode("utf-8"))
    if extracted_bytes > remaining_budget_bytes:
        return await create_failed_uploaded_file_source(
            db,
            user_id=user_id,
            agent_id=agent_id,
            filename=filename,
            content_type=content_type,
            uploaded_bytes=len(payload),
            error_message=(
                f"File exceeds remaining storage budget ({extracted_bytes}B needed, "
                f"{remaining_budget_bytes}B available)."
            ),
        )

    title = (filename or "Uploaded file").strip()[:255] or "Uploaded file"
    source = await create_source(
        db,
        user_id,
        KnowledgeSourceCreateRequest(
            agent_id=agent_id,
            type="file",
            title=title,
            storage_bucket="inline-upload",
            storage_path=_safe_storage_path(filename),
            metadata={
                "origin": "dashboard_file_upload",
                "filename": filename,
                "content_type": content_type,
                "uploaded_bytes": len(payload),
                "extracted_text": extracted,
            },
        ),
    )
    try:
        source_out, job = await index_file_source(db, source.id, user_id)
        return FileUploadResultDTO(source=source_out, job=job, status="succeeded", error_message=None)
    except AppError:
        failed = await _load_source(db, source.id, user_id)
        return FileUploadResultDTO(source=failed, job=None, status="failed", error_message=failed.error_message)


async def get_jobs(db: AsyncSession, source_id: UUID, user_id: UUID) -> list[IndexJobDTO]:
    result = await db.execute(
        text(
            """
            select
              id, knowledge_source_id, agent_id, user_id, status, attempt, triggered_by, error_message,
              started_at, finished_at, phase, pages_total, pages_processed, chunks_total, chunks_embedded, progress_pct, metrics, created_at, updated_at
            from public.indexing_jobs
            where knowledge_source_id = :source_id and user_id = :user_id
            order by created_at desc
            """
        ),
        {"source_id": str(source_id), "user_id": str(user_id)},
    )
    return [IndexJobDTO.model_validate(row) for row in result.mappings().all()]


async def get_latest_job(db: AsyncSession, source_id: UUID, user_id: UUID) -> IndexJobDTO | None:
    result = await db.execute(
        text(
            """
            select
              id, knowledge_source_id, agent_id, user_id, status, attempt, triggered_by, error_message,
              started_at, finished_at, phase, pages_total, pages_processed, chunks_total, chunks_embedded, progress_pct, metrics, created_at, updated_at
            from public.indexing_jobs
            where knowledge_source_id = :source_id and user_id = :user_id
            order by created_at desc
            limit 1
            """
        ),
        {"source_id": str(source_id), "user_id": str(user_id)},
    )
    row = result.mappings().first()
    return IndexJobDTO.model_validate(row) if row else None


async def list_website_sources_for_agent(
    db: AsyncSession, user_id: UUID, agent_id: UUID
) -> list[WebsiteSourceListItemDTO]:
    agent_check = await db.execute(
        text("select id from public.agents where id = :agent_id and user_id = :user_id"),
        {"agent_id": str(agent_id), "user_id": str(user_id)},
    )
    if agent_check.mappings().first() is None:
        raise AppError(code="agents.not_found", message="Agent not found", status_code=404)
    _, _, plan_features = await _fetch_active_subscription_plan(db, user_id)
    storage_cap_bytes = _included_storage_bytes_from_plan_features(plan_features)

    result = await db.execute(
        text(
            """
            select
              s.id,
              s.agent_id,
              s.title,
              s.source_url,
              s.metadata,
              s.status::text as status,
              s.last_indexed_at,
              s.error_message as error_message,
              coalesce(cnt.c, 0) as link_count,
              j.status::text as latest_job_status,
              j.phase::text as latest_job_phase,
              j.pages_total as job_pages_total,
              j.pages_processed as job_pages_processed,
              j.progress_pct as job_progress_pct,
              j.metrics as job_metrics,
              (s.metadata->>'website_mode') as website_mode
            from public.knowledge_sources s
            left join lateral (
              select count(*)::int as c
              from public.knowledge_source_pages p
              where p.knowledge_source_id = s.id
                and p.status <> 'excluded'::public.crawl_page_status
            ) cnt on true
            left join lateral (
              select status, phase, pages_total, pages_processed, progress_pct, metrics
              from public.indexing_jobs j2
              where j2.knowledge_source_id = s.id
              order by j2.created_at desc
              limit 1
            ) j on true
            where s.user_id = :user_id
              and s.agent_id = :agent_id
              and s.type = 'website'
            order by s.created_at desc
            """
        ),
        {"user_id": str(user_id), "agent_id": str(agent_id)},
    )
    out: list[WebsiteSourceListItemDTO] = []
    for row in result.mappings().all():
        wm = row.get("website_mode")
        mode = wm if wm in ("crawl", "sitemap", "individual") else None
        lim = _job_crawl_limit_exceeded(
            str(row["latest_job_status"]) if row.get("latest_job_status") else None,
            row.get("job_metrics"),
            row.get("job_pages_total"),
            row.get("job_pages_processed"),
            storage_cap_bytes=storage_cap_bytes,
        )
        meta_raw = row.get("metadata")
        meta = dict(meta_raw) if isinstance(meta_raw, dict) else {}
        dup_r = meta.get("duplicate_reason")
        dup_sid = meta.get("duplicate_of_source_id")
        jm_raw = row.get("job_metrics")
        job_metrics_dict = dict(jm_raw) if isinstance(jm_raw, dict) else None
        out.append(
            WebsiteSourceListItemDTO(
                id=row["id"],
                agent_id=row["agent_id"],
                title=str(row["title"]),
                source_url=str(row["source_url"]) if row["source_url"] else None,
                status=str(row["status"]),
                website_mode=mode,  # type: ignore[arg-type]
                link_count=int(row["link_count"] or 0),
                last_indexed_at=row["last_indexed_at"],
                error_message=str(row["error_message"]).strip() if row.get("error_message") else None,
                latest_job_status=str(row["latest_job_status"]) if row["latest_job_status"] else None,
                latest_job_phase=str(row["latest_job_phase"]) if row["latest_job_phase"] else None,
                job_metrics=job_metrics_dict,
                job_pages_total=int(row["job_pages_total"]) if row.get("job_pages_total") is not None else None,
                job_pages_processed=int(row["job_pages_processed"]) if row.get("job_pages_processed") is not None else None,
                job_progress_pct=int(row["job_progress_pct"]) if row.get("job_progress_pct") is not None else None,
                job_crawl_limit_exceeded=lim,
                reindexed_duplicate=bool(meta.get("reindexed_duplicate", False)),
                duplicate_reason=str(dup_r) if dup_r else None,
                duplicate_of_source_id=str(dup_sid) if dup_sid else None,
            )
        )
    return out


async def list_file_sources_for_agent(
    db: AsyncSession, user_id: UUID, agent_id: UUID
) -> list[FileSourceListItemDTO]:
    agent_check = await db.execute(
        text("select id from public.agents where id = :agent_id and user_id = :user_id"),
        {"agent_id": str(agent_id), "user_id": str(user_id)},
    )
    if agent_check.mappings().first() is None:
        raise AppError(code="agents.not_found", message="Agent not found", status_code=404)

    result = await db.execute(
        text(
            """
            select
              s.id,
              s.agent_id,
              s.title,
              s.storage_bucket,
              s.storage_path,
              s.status::text as status,
              s.last_indexed_at,
              coalesce(chars.character_count, 0) as character_count,
              j.status::text as latest_job_status,
              j.phase::text as latest_job_phase,
              j.progress_pct as job_progress_pct
            from public.knowledge_sources s
            left join lateral (
              select coalesce(sum(char_length(c.content))::bigint, 0) as character_count
              from public.knowledge_chunks c
              where c.knowledge_source_id = s.id
            ) chars on true
            left join lateral (
              select status, phase, progress_pct
              from public.indexing_jobs j2
              where j2.knowledge_source_id = s.id
              order by j2.created_at desc
              limit 1
            ) j on true
            where s.user_id = :user_id
              and s.agent_id = :agent_id
              and s.type = 'file'
            order by s.created_at desc
            """
        ),
        {"user_id": str(user_id), "agent_id": str(agent_id)},
    )
    return [FileSourceListItemDTO.model_validate(row) for row in result.mappings().all()]


async def list_website_source_pages(
    db: AsyncSession, user_id: UUID, source_id: UUID, *, offset: int, limit: int
) -> tuple[list[dict[str, object]], int]:
    src = (
        await db.execute(
            text(
                """
                select id
                from public.knowledge_sources
                where id = :source_id and user_id = :user_id and type = 'website'
                """
            ),
            {"source_id": str(source_id), "user_id": str(user_id)},
        )
    ).first()
    if src is None:
        raise AppError(code="knowledge.source_not_found", message="Website source not found", status_code=404)

    total_row = (
        await db.execute(
            text(
                """
                select count(*)::int as c
                from public.knowledge_source_pages p
                where p.knowledge_source_id = :source_id and p.user_id = :user_id
                """
            ),
            {"source_id": str(source_id), "user_id": str(user_id)},
        )
    ).mappings().one()
    total = int(total_row["c"] or 0)

    rows = (
        await db.execute(
            text(
                """
                select p.id, p.url, p.status::text as status, p.depth, p.last_crawled_at as last_indexed_at, p.http_status
                from public.knowledge_source_pages p
                where p.knowledge_source_id = :source_id and p.user_id = :user_id
                order by p.url asc
                limit :limit offset :offset
                """
            ),
            {"source_id": str(source_id), "user_id": str(user_id), "limit": limit, "offset": offset},
        )
    ).mappings().all()
    pages = [dict(r) for r in rows]
    return pages, total


async def delete_website_source(db: AsyncSession, user_id: UUID, source_id: UUID) -> None:
    result = await db.execute(
        text(
            """
            delete from public.knowledge_sources
            where id = :source_id and user_id = :user_id and type = 'website'
            returning id
            """
        ),
        {"source_id": str(source_id), "user_id": str(user_id)},
    )
    if result.first() is None:
        raise AppError(code="knowledge.source_not_found", message="Website source not found", status_code=404)
    await db.commit()


async def delete_website_source_page(
    db: AsyncSession, user_id: UUID, source_id: UUID, page_id: UUID
) -> None:
    """Remove a single crawled page row + its associated chunks from the parent website source."""
    page = (
        await db.execute(
            text(
                """
                select p.id, p.url
                from public.knowledge_source_pages p
                join public.knowledge_sources s on s.id = p.knowledge_source_id
                where p.id = :page_id
                  and p.knowledge_source_id = :source_id
                  and s.user_id = :user_id
                  and s.type = 'website'
                """
            ),
            {"page_id": str(page_id), "source_id": str(source_id), "user_id": str(user_id)},
        )
    ).mappings().first()
    if page is None:
        raise AppError(code="knowledge.page_not_found", message="Website page not found", status_code=404)
    await db.execute(
        text(
            """
            delete from public.knowledge_chunks
            where knowledge_source_id = :source_id
              and metadata->>'page_url' = :url
            """
        ),
        {"source_id": str(source_id), "url": str(page["url"])},
    )
    await db.execute(
        text("delete from public.knowledge_source_pages where id = :page_id"),
        {"page_id": str(page_id)},
    )
    await db.commit()


async def update_website_source_page(
    db: AsyncSession,
    user_id: UUID,
    source_id: UUID,
    page_id: UUID,
    new_url: str,
) -> None:
    """Replace a single page's URL/content: refetch, re-embed, swap in new chunks for the page."""
    page = (
        await db.execute(
            text(
                """
                select p.id, p.url, p.knowledge_source_id, s.agent_id, s.source_url
                from public.knowledge_source_pages p
                join public.knowledge_sources s on s.id = p.knowledge_source_id
                where p.id = :page_id
                  and p.knowledge_source_id = :source_id
                  and s.user_id = :user_id
                  and s.type = 'website'
                """
            ),
            {"page_id": str(page_id), "source_id": str(source_id), "user_id": str(user_id)},
        )
    ).mappings().first()
    if page is None:
        raise AppError(code="knowledge.page_not_found", message="Website page not found", status_code=404)

    cleaned_url = (new_url or "").strip()
    if not cleaned_url:
        raise AppError(code="knowledge.invalid_url", message="URL cannot be empty", status_code=422)

    old_url = str(page["url"])
    agent_id = UUID(str(page["agent_id"]))
    source_url = page["source_url"] if page["source_url"] is not None else None

    if cleaned_url != old_url:
        conflict = (
            await db.execute(
                text(
                    """
                    select id
                    from public.knowledge_source_pages
                    where knowledge_source_id = :sid and url = :url and id <> :pid
                    """
                ),
                {"sid": str(source_id), "url": cleaned_url, "pid": str(page_id)},
            )
        ).mappings().first()
        if conflict is not None:
            raise AppError(
                code="knowledge.page_url_conflict",
                message="Another page with this URL already exists in this source",
                status_code=409,
            )

    await db.execute(
        text(
            """
            delete from public.knowledge_chunks
            where knowledge_source_id = :source_id
              and metadata->>'page_url' = :url
            """
        ),
        {"source_id": str(source_id), "url": old_url},
    )

    await db.execute(
        text(
            """
            update public.knowledge_source_pages
            set url = :url,
                status = 'queued'::public.crawl_page_status,
                http_status = null,
                last_crawled_at = null
            where id = :page_id
            """
        ),
        {"url": cleaned_url, "page_id": str(page_id)},
    )
    await db.commit()

    pages, _, _, _, _ = await _fetch_pages_for_urls([cleaned_url], crawl_budget_bytes=None)
    page_data = pages[0] if pages else {"text": "", "title": None, "http_status": None}
    page_text = str(page_data.get("text") or "")
    page_title = str(page_data.get("title") or "").strip()
    http_status = page_data.get("http_status")

    if not page_text.strip():
        await db.execute(
            text(
                """
                update public.knowledge_source_pages
                set status = 'failed'::public.crawl_page_status,
                    http_status = :http,
                    last_crawled_at = now()
                where id = :page_id
                """
            ),
            {"page_id": str(page_id), "http": http_status},
        )
        await db.commit()
        return

    prefix_parts = [p for p in [page_title, cleaned_url] if p]
    chunk_prefix = " | ".join(prefix_parts).strip()
    chunk_records: list[dict[str, object]] = []
    for page_chunk_index, chunk in enumerate(_chunk_text(page_text)):
        chunk_with_context = f"{chunk_prefix}\n\n{chunk}".strip() if chunk_prefix else chunk
        chunk_records.append(
            {
                "content": chunk_with_context,
                "raw_content": chunk,
                "page_chunk_index": page_chunk_index,
            }
        )

    if not chunk_records:
        await db.execute(
            text(
                """
                update public.knowledge_source_pages
                set status = 'parsed'::public.crawl_page_status,
                    http_status = :http,
                    last_crawled_at = now()
                where id = :page_id
                """
            ),
            {"page_id": str(page_id), "http": http_status},
        )
        await db.commit()
        return

    embeddings = await _embed_texts([str(rec["content"]) for rec in chunk_records])

    max_idx_row = (
        await db.execute(
            text(
                "select coalesce(max(chunk_index), -1) as m from public.knowledge_chunks where knowledge_source_id = :sid"
            ),
            {"sid": str(source_id)},
        )
    ).mappings().first()
    base_idx = int(max_idx_row["m"]) + 1 if max_idx_row else 0

    for offset, (rec, embedding) in enumerate(zip(chunk_records, embeddings)):
        if len(embedding) != EMBEDDING_DIMENSION:
            raise AppError(
                code="knowledge.embedding_dimension_mismatch",
                message="Indexed content format does not match storage",
                status_code=500,
                details={"expected": EMBEDDING_DIMENSION, "actual": len(embedding)},
            )
        chunk = str(rec["content"])
        idx = base_idx + offset
        await db.execute(
            text(
                """
                insert into public.knowledge_chunks (
                  agent_id, user_id, knowledge_source_id, chunk_index, content, embedding, token_count, metadata
                ) values (
                  :agent_id, :user_id, :knowledge_source_id, :chunk_index, :content, CAST(:embedding AS vector), :token_count, CAST(:metadata AS jsonb)
                )
                """
            ),
            {
                "agent_id": str(agent_id),
                "user_id": str(user_id),
                "knowledge_source_id": str(source_id),
                "chunk_index": idx,
                "content": chunk,
                "embedding": _vector_literal(embedding),
                "token_count": _token_estimate(chunk),
                "metadata": json.dumps(
                    {
                        "source_url": source_url,
                        "page_url": cleaned_url,
                        "page_title": page_title,
                        "raw_content": rec.get("raw_content"),
                        "chunk_index": idx,
                        "page_chunk_index": rec.get("page_chunk_index"),
                    }
                ),
            },
        )

    await db.execute(
        text(
            """
            update public.knowledge_source_pages
            set status = 'parsed'::public.crawl_page_status,
                http_status = :http,
                last_crawled_at = now()
            where id = :page_id
            """
        ),
        {"page_id": str(page_id), "http": http_status},
    )
    await db.execute(
        text(
            """
            update public.knowledge_sources
            set last_indexed_at = now()
            where id = :sid
            """
        ),
        {"sid": str(source_id)},
    )
    await db.commit()


async def delete_file_source(db: AsyncSession, user_id: UUID, source_id: UUID) -> None:
    result = await db.execute(
        text(
            """
            delete from public.knowledge_sources
            where id = :source_id and user_id = :user_id and type = 'file'
            returning id
            """
        ),
        {"source_id": str(source_id), "user_id": str(user_id)},
    )
    if result.first() is None:
        raise AppError(code="knowledge.source_not_found", message="File source not found", status_code=404)
    await db.commit()


async def list_text_snippet_sources_for_agent(
    db: AsyncSession, user_id: UUID, agent_id: UUID
) -> list[TextSnippetListItemDTO]:
    agent_check = await db.execute(
        text("select id from public.agents where id = :agent_id and user_id = :user_id"),
        {"agent_id": str(agent_id), "user_id": str(user_id)},
    )
    if agent_check.mappings().first() is None:
        raise AppError(code="agents.not_found", message="Agent not found", status_code=404)

    result = await db.execute(
        text(
            """
            select
              s.id,
              s.agent_id,
              s.title,
              s.status::text as status,
              s.last_indexed_at,
              s.updated_at,
              coalesce(chars.character_count, 0) as character_count,
              coalesce(left(s.raw_text, 400), '') as preview,
              j.status::text as latest_job_status,
              j.phase::text as latest_job_phase,
              j.progress_pct as job_progress_pct
            from public.knowledge_sources s
            left join lateral (
              select coalesce(sum(char_length(c.content))::bigint, 0) as character_count
              from public.knowledge_chunks c
              where c.knowledge_source_id = s.id
            ) chars on true
            left join lateral (
              select status, phase, progress_pct
              from public.indexing_jobs j2
              where j2.knowledge_source_id = s.id
              order by j2.created_at desc
              limit 1
            ) j on true
            where s.user_id = :user_id
              and s.agent_id = :agent_id
              and s.type = 'text_snippet'
            order by s.updated_at desc
            """
        ),
        {"user_id": str(user_id), "agent_id": str(agent_id)},
    )
    return [TextSnippetListItemDTO.model_validate(row) for row in result.mappings().all()]


async def get_text_snippet_detail(db: AsyncSession, user_id: UUID, source_id: UUID) -> TextSnippetDetailDTO:
    row = (
        await db.execute(
            text(
                """
                select
                  id,
                  agent_id,
                  title,
                  coalesce(raw_text, '') as text,
                  status::text as status,
                  last_indexed_at,
                  updated_at
                from public.knowledge_sources
                where id = :source_id and user_id = :user_id and type = 'text_snippet'
                """
            ),
            {"source_id": str(source_id), "user_id": str(user_id)},
        )
    ).mappings().first()
    if row is None:
        raise AppError(code="knowledge.source_not_found", message="Text snippet not found", status_code=404)
    return TextSnippetDetailDTO(
        id=row["id"],
        agent_id=row["agent_id"],
        title=str(row["title"]),
        text=str(row["text"]),
        status=str(row["status"]),
        last_indexed_at=row["last_indexed_at"],
        updated_at=row["updated_at"],
    )


async def update_text_snippet_source(
    db: AsyncSession, user_id: UUID, source_id: UUID, *, title: str, snippet_text: str
) -> tuple[KnowledgeSourceDTO, IndexJobDTO]:
    src = await _load_source(db, source_id, user_id)
    if src.type != "text_snippet":
        raise AppError(code="validation.invalid_input", message="Only text_snippet sources can be updated here", status_code=422)
    body = (snippet_text or "").strip()
    t = (title or "").strip()[:255]
    if not body or not t:
        raise AppError(code="validation.invalid_input", message="Title and text are required", status_code=422)

    _, _, plan_features = await _fetch_active_subscription_plan(db, user_id)
    included_storage_cap_bytes = _included_storage_bytes_from_plan_features(plan_features)
    effective_storage_cap_bytes = _effective_storage_cap_bytes(included_storage_cap_bytes)
    currently_used_storage_bytes = await _agent_used_storage_bytes(db, user_id=user_id, agent_id=src.agent_id)
    remaining_budget_bytes = max(0, effective_storage_cap_bytes - currently_used_storage_bytes)
    if remaining_budget_bytes <= 0:
        raise AppError(
            code="knowledge.storage_budget_exhausted",
            message="Knowledge storage limit reached. Remove or delete indexed content, or upgrade your plan.",
            status_code=422,
        )

    await db.execute(
        text(
            """
            update public.knowledge_sources
            set title = :title, raw_text = :raw_text, status = 'pending', error_message = null, updated_at = now()
            where id = :source_id and user_id = :user_id and type = 'text_snippet'
            """
        ),
        {"title": t, "raw_text": body, "source_id": str(source_id), "user_id": str(user_id)},
    )
    await db.commit()
    return await index_text_snippet_source(db, source_id, user_id)


async def delete_text_snippet_source(db: AsyncSession, user_id: UUID, source_id: UUID) -> None:
    result = await db.execute(
        text(
            """
            delete from public.knowledge_sources
            where id = :source_id and user_id = :user_id and type = 'text_snippet'
            returning id
            """
        ),
        {"source_id": str(source_id), "user_id": str(user_id)},
    )
    if result.first() is None:
        raise AppError(code="knowledge.source_not_found", message="Text snippet not found", status_code=404)
    await db.commit()


async def list_qa_sources_for_agent(db: AsyncSession, user_id: UUID, agent_id: UUID) -> list[QAPairListItemDTO]:
    agent_check = await db.execute(
        text("select id from public.agents where id = :agent_id and user_id = :user_id"),
        {"agent_id": str(agent_id), "user_id": str(user_id)},
    )
    if agent_check.mappings().first() is None:
        raise AppError(code="agents.not_found", message="Agent not found", status_code=404)

    result = await db.execute(
        text(
            """
            select
              s.id,
              s.agent_id,
              s.title,
              s.status::text as status,
              s.last_indexed_at,
              s.updated_at,
              q.question,
              coalesce(left(q.answer, 400), '') as answer_preview,
              coalesce(chars.character_count, 0) as character_count,
              j.status::text as latest_job_status,
              j.phase::text as latest_job_phase,
              j.progress_pct as job_progress_pct
            from public.knowledge_sources s
            inner join public.knowledge_qa_items q on q.knowledge_source_id = s.id
            left join lateral (
              select coalesce(sum(char_length(c.content))::bigint, 0) as character_count
              from public.knowledge_chunks c
              where c.knowledge_source_id = s.id
            ) chars on true
            left join lateral (
              select status, phase, progress_pct
              from public.indexing_jobs j2
              where j2.knowledge_source_id = s.id
              order by j2.created_at desc
              limit 1
            ) j on true
            where s.user_id = :user_id
              and s.agent_id = :agent_id
              and s.type = 'q_and_a'
            order by s.updated_at desc
            """
        ),
        {"user_id": str(user_id), "agent_id": str(agent_id)},
    )
    return [QAPairListItemDTO.model_validate(row) for row in result.mappings().all()]


async def get_qa_pair_detail(db: AsyncSession, user_id: UUID, source_id: UUID) -> QAPairDetailDTO:
    row = (
        await db.execute(
            text(
                """
                select
                  s.id,
                  s.agent_id,
                  q.question,
                  q.answer,
                  s.status::text as status,
                  s.last_indexed_at,
                  s.updated_at
                from public.knowledge_sources s
                inner join public.knowledge_qa_items q on q.knowledge_source_id = s.id
                where s.id = :source_id and s.user_id = :user_id and s.type = 'q_and_a'
                order by q.created_at asc
                limit 1
                """
            ),
            {"source_id": str(source_id), "user_id": str(user_id)},
        )
    ).mappings().first()
    if row is None:
        raise AppError(code="knowledge.source_not_found", message="Q&A pair not found", status_code=404)
    return QAPairDetailDTO(
        id=row["id"],
        agent_id=row["agent_id"],
        question=str(row["question"]),
        answer=str(row["answer"]),
        status=str(row["status"]),
        last_indexed_at=row["last_indexed_at"],
        updated_at=row["updated_at"],
    )


async def update_qa_pair_source(
    db: AsyncSession, user_id: UUID, source_id: UUID, *, question: str, answer: str
) -> tuple[KnowledgeSourceDTO, IndexJobDTO]:
    src = await _load_source(db, source_id, user_id)
    if src.type != "q_and_a":
        raise AppError(code="validation.invalid_input", message="Only q_and_a sources can be updated here", status_code=422)
    q = (question or "").strip()
    a = (answer or "").strip()
    if not q or not a:
        raise AppError(code="validation.invalid_input", message="Question and answer are required", status_code=422)
    title = q[:255]

    _, _, plan_features = await _fetch_active_subscription_plan(db, user_id)
    included_storage_cap_bytes = _included_storage_bytes_from_plan_features(plan_features)
    effective_storage_cap_bytes = _effective_storage_cap_bytes(included_storage_cap_bytes)
    currently_used_storage_bytes = await _agent_used_storage_bytes(db, user_id=user_id, agent_id=src.agent_id)
    remaining_budget_bytes = max(0, effective_storage_cap_bytes - currently_used_storage_bytes)
    if remaining_budget_bytes <= 0:
        raise AppError(
            code="knowledge.storage_budget_exhausted",
            message="Knowledge storage limit reached. Remove or delete indexed content, or upgrade your plan.",
            status_code=422,
        )

    upd = await db.execute(
        text(
            """
            update public.knowledge_qa_items q
            set question = :question, answer = :answer, updated_at = now()
            from public.knowledge_sources s
            where q.knowledge_source_id = s.id
              and s.id = :source_id and s.user_id = :user_id and s.type = 'q_and_a'
            returning q.id
            """
        ),
        {"question": q, "answer": a, "source_id": str(source_id), "user_id": str(user_id)},
    )
    if upd.first() is None:
        raise AppError(code="knowledge.source_not_found", message="Q&A pair not found", status_code=404)

    await db.execute(
        text(
            """
            update public.knowledge_sources
            set title = :title, status = 'pending', error_message = null, updated_at = now()
            where id = :source_id and user_id = :user_id and type = 'q_and_a'
            """
        ),
        {"title": title, "source_id": str(source_id), "user_id": str(user_id)},
    )
    await db.commit()
    return await index_qa_source(db, source_id, user_id)


async def delete_qa_source(db: AsyncSession, user_id: UUID, source_id: UUID) -> None:
    result = await db.execute(
        text(
            """
            delete from public.knowledge_sources
            where id = :source_id and user_id = :user_id and type = 'q_and_a'
            returning id
            """
        ),
        {"source_id": str(source_id), "user_id": str(user_id)},
    )
    if result.first() is None:
        raise AppError(code="knowledge.source_not_found", message="Q&A pair not found", status_code=404)
    await db.commit()


def _included_storage_bytes_from_plan_features(features: dict[str, object]) -> int:
    """Total indexed knowledge cap (website + files + snippets + Q&A share one pool)."""
    if not isinstance(features, dict):
        return DEFAULT_KNOWLEDGE_STORAGE_CAP_BYTES
    return training_storage_cap_bytes(features)


def _effective_storage_cap_bytes(included_storage_bytes: int) -> int:
    """
    Internal allowance only: hobby/free 500KB-style plans receive hidden +23% storage headroom.
    """
    if included_storage_bytes == STARTER_KNOWLEDGE_STORAGE_CAP_BYTES:
        return int(included_storage_bytes * (1 + STARTER_HIDDEN_STORAGE_GRACE_RATIO))
    return included_storage_bytes


async def _agent_used_storage_bytes(db: AsyncSession, *, user_id: UUID, agent_id: UUID) -> int:
    usage_row = (
        await db.execute(
            text(
                """
                select coalesce(sum(
                  octet_length(c.content)
                )::bigint, 0) as used_bytes
                from public.knowledge_chunks c
                join public.knowledge_sources s on s.id = c.knowledge_source_id
                where s.agent_id = :agent_id and s.user_id = :user_id
                """
            ),
            {"agent_id": str(agent_id), "user_id": str(user_id)},
        )
    ).mappings().one()
    return int(usage_row["used_bytes"] or 0)


def _coerce_job_metrics(metrics: object) -> dict[str, Any]:
    if metrics is None:
        return {}
    if isinstance(metrics, dict):
        return dict(metrics)
    if isinstance(metrics, str):
        try:
            return dict(json.loads(metrics))
        except json.JSONDecodeError:
            return {}
    return {}


def _job_crawl_limit_exceeded(
    job_status: str | None,
    metrics: object,
    pages_total: object,
    pages_processed: object,
    storage_cap_bytes: int | None = None,
) -> bool:
    if (job_status or "").lower() != "succeeded":
        return False
    m = _coerce_job_metrics(metrics)
    if str(m.get("crawl_stopped_reason")) != "budget":
        return False
    if storage_cap_bytes is not None and storage_cap_bytes > 0:
        try:
            indexed_bytes = int(m.get("indexed_source_bytes") or 0)
        except (TypeError, ValueError):
            indexed_bytes = 0
        if indexed_bytes < int(storage_cap_bytes * 0.95):
            return False
    try:
        pt = int(pages_total) if pages_total is not None else None
        pp = int(pages_processed) if pages_processed is not None else None
    except (TypeError, ValueError):
        return False
    return pt is not None and pp is not None and pp < pt


def _url_duplicate_key(url: str) -> str:
    """Lowercase URL with trailing slashes removed (for duplicate detection)."""
    return re.sub(r"/+$", "", url.strip().lower())


def _canonical_path_rule_pairs(rules: list[dict[str, str]]) -> list[tuple[str, str]]:
    pairs = sorted(
        ((str(r.get("operator", "")), str(r.get("pattern", ""))) for r in rules),
        key=lambda t: (t[0], t[1]),
    )
    return pairs


def _dashboard_path_rules_match(
    incoming_include: list[dict[str, str]],
    incoming_exclude: list[dict[str, str]],
    stored_metadata: dict[str, Any] | None,
) -> bool:
    """True when stored source metadata has the same include/exclude path rules (order-independent)."""
    md = stored_metadata if isinstance(stored_metadata, dict) else {}
    raw_inc = md.get("include_rules") or []
    raw_exc = md.get("exclude_rules") or []
    if not isinstance(raw_inc, list):
        raw_inc = []
    if not isinstance(raw_exc, list):
        raw_exc = []
    st_inc = [dict(x) for x in raw_inc if isinstance(x, dict)]
    st_exc = [dict(x) for x in raw_exc if isinstance(x, dict)]
    return _canonical_path_rule_pairs(incoming_include) == _canonical_path_rule_pairs(
        st_inc
    ) and _canonical_path_rule_pairs(incoming_exclude) == _canonical_path_rule_pairs(st_exc)


def _website_filter_summary(
    include_rules: list[dict[str, str]], exclude_rules: list[dict[str, str]]
) -> dict[str, list[dict[str, str]]]:
    """Canonical filter lists for job metrics (order-independent)."""
    inc = [{"operator": o, "pattern": p} for o, p in _canonical_path_rule_pairs(include_rules)]
    exc = [{"operator": o, "pattern": p} for o, p in _canonical_path_rule_pairs(exclude_rules)]
    return {"include_rules": inc, "exclude_rules": exc}


def _website_fetch_page_stats(pages: list[dict[str, object]]) -> dict[str, int]:
    """Counts from in-memory page dicts after HTTP fetch (before embedding)."""
    total = len(pages)
    empty_text = sum(1 for p in pages if not str(p.get("text") or "").strip())
    http_missing = sum(1 for p in pages if p.get("http_status") is None)
    http_4xx = sum(
        1
        for p in pages
        if isinstance(p.get("http_status"), int) and 400 <= int(p["http_status"]) < 500
    )
    http_5xx = sum(
        1
        for p in pages
        if isinstance(p.get("http_status"), int) and int(p["http_status"]) >= 500
    )
    return {
        "urls_fetched": total,
        "urls_empty_text": empty_text,
        "urls_with_text": total - empty_text,
        "urls_http_missing": http_missing,
        "urls_http_4xx": http_4xx,
        "urls_http_5xx": http_5xx,
    }


async def _fetch_active_subscription_plan(
    db: AsyncSession, user_id: UUID
) -> tuple[str, str, dict[str, Any]]:
    plan_row = (
        await db.execute(
            text(
                """
                select p.slug::text as slug, p.name::text as name, p.features
                from public.subscriptions s
                join public.plans p on p.id = s.plan_id
                where s.user_id = :user_id
                  and s.status in ('trialing', 'active', 'past_due')
                order by s.current_period_end desc
                limit 1
                """
            ),
            {"user_id": str(user_id)},
        )
    ).mappings().first()
    if plan_row is None:
        return "free", "Free", {}
    fr = plan_row["features"]
    features: dict[str, Any] = dict(fr) if isinstance(fr, dict) else {}
    return str(plan_row["slug"]), str(plan_row["name"]), features


async def _find_dashboard_website_duplicate(
    db: AsyncSession,
    user_id: UUID,
    agent_id: UUID,
    normalized_url: str,
    *,
    include_rules: list[dict[str, str]],
    exclude_rules: list[dict[str, str]],
) -> tuple[UUID, str] | None:
    """If the same seed URL + path rules already exist, return (source_id, reason).

    Path include/exclude lists are part of identity: a filtered crawl (e.g. ``contains /foo/``)
    is not a duplicate of an unfiltered crawl that indexed the same homepage URL.
    """
    key = _url_duplicate_key(normalized_url)
    r_pages = (
        await db.execute(
            text(
                """
                select s.id::text as sid, s.metadata
                from public.knowledge_source_pages p
                join public.knowledge_sources s on s.id = p.knowledge_source_id
                where s.user_id = cast(:user_id as uuid)
                  and s.agent_id = cast(:agent_id as uuid)
                  and s.type = 'website'
                  and s.status::text <> 'skipped_duplicate'
                  and regexp_replace(lower(btrim(p.url)), '/+$', '') = :url_key
                """
            ),
            {"user_id": str(user_id), "agent_id": str(agent_id), "url_key": key},
        )
    ).mappings().all()
    for row in r_pages:
        md = row.get("metadata")
        meta_dict = dict(md) if isinstance(md, dict) else {}
        if _dashboard_path_rules_match(include_rules, exclude_rules, meta_dict):
            return UUID(str(row["sid"])), "page_already_indexed"

    r_roots = (
        await db.execute(
            text(
                """
                select id::text as sid, metadata
                from public.knowledge_sources
                where user_id = cast(:user_id as uuid)
                  and agent_id = cast(:agent_id as uuid)
                  and type = 'website'
                  and status::text not in ('skipped_duplicate', 'failed')
                  and source_url is not null
                  and regexp_replace(lower(btrim(source_url)), '/+$', '') = :url_key
                """
            ),
            {"user_id": str(user_id), "agent_id": str(agent_id), "url_key": key},
        )
    ).mappings().all()
    for row in r_roots:
        md = row.get("metadata")
        meta_dict = dict(md) if isinstance(md, dict) else {}
        if _dashboard_path_rules_match(include_rules, exclude_rules, meta_dict):
            return UUID(str(row["sid"])), "same_root_url"
    return None


async def get_agent_website_usage(db: AsyncSession, user_id: UUID, agent_id: UUID) -> WebsiteUsageResponse:
    agent_check = await db.execute(
        text("select id from public.agents where id = :agent_id and user_id = :user_id"),
        {"agent_id": str(agent_id), "user_id": str(user_id)},
    )
    if agent_check.mappings().first() is None:
        raise AppError(code="agents.not_found", message="Agent not found", status_code=404)

    usage_row = (
        await db.execute(
            text(
                """
                with chunk_bytes as (
                  select
                    s.type::text as type,
                    coalesce(sum(
                      octet_length(c.content)
                    )::bigint, 0) as bytes
                  from public.knowledge_sources s
                  left join public.knowledge_chunks c on c.knowledge_source_id = s.id
                  where s.agent_id = :agent_id and s.user_id = :user_id
                  group by s.type
                )
                select
                  coalesce((
                    select count(*)::bigint
                    from public.knowledge_sources s
                    where s.agent_id = :agent_id and s.user_id = :user_id and s.type = 'file'
                  ), 0) as total_files,
                  coalesce((
                    select count(*)::bigint
                    from public.knowledge_sources s
                    where s.agent_id = :agent_id and s.user_id = :user_id and s.type = 'text_snippet'
                  ), 0) as total_snippets,
                  coalesce((
                    select count(*)::bigint
                    from public.knowledge_sources s
                    where s.agent_id = :agent_id and s.user_id = :user_id and s.type = 'q_and_a'
                  ), 0) as total_qa_pairs,
                  coalesce((
                    select count(*)::bigint
                    from public.knowledge_source_pages p
                    join public.knowledge_sources s on s.id = p.knowledge_source_id
                    where s.agent_id = :agent_id and s.user_id = :user_id and s.type = 'website'
                      and p.status in (
                        'parsed'::public.crawl_page_status,
                        'failed'::public.crawl_page_status,
                        'fetched'::public.crawl_page_status
                      )
                  ), 0) as total_links,
                  coalesce((select sum(bytes) from chunk_bytes), 0) as used_bytes,
                  coalesce((select bytes from chunk_bytes where type = 'website'), 0) as website_used_bytes,
                  coalesce((select bytes from chunk_bytes where type = 'file'), 0) as files_used_bytes,
                  coalesce((select bytes from chunk_bytes where type = 'text_snippet'), 0) as snippets_used_bytes,
                  coalesce((select bytes from chunk_bytes where type = 'q_and_a'), 0) as qa_used_bytes
                """
            ),
            {"agent_id": str(agent_id), "user_id": str(user_id)},
        )
    ).mappings().one()

    plan_slug, plan_name, features = await _fetch_active_subscription_plan(db, user_id)

    included = _included_storage_bytes_from_plan_features(features)
    effective_storage_cap_bytes = _effective_storage_cap_bytes(included)
    used = int(usage_row["used_bytes"] or 0)
    total_links = int(usage_row["total_links"] or 0)
    total_files = int(usage_row["total_files"] or 0)
    total_snippets = int(usage_row["total_snippets"] or 0)
    total_qa_pairs = int(usage_row["total_qa_pairs"] or 0)
    website_used = int(usage_row.get("website_used_bytes") or 0)
    files_used = int(usage_row.get("files_used_bytes") or 0)
    snippets_used = int(usage_row.get("snippets_used_bytes") or 0)
    qa_used = int(usage_row.get("qa_used_bytes") or 0)
    show_upgrade = plan_slug == "free" or used > included
    remaining_total = max(0, int(effective_storage_cap_bytes) - used)

    running_crawl_row = (
        await db.execute(
            text(
                """
                select (j.metrics->>'crawl_http_bytes')::bigint as b
                from public.indexing_jobs j
                join public.knowledge_sources s on s.id = j.knowledge_source_id
                where s.agent_id = cast(:agent_id as uuid)
                  and s.user_id = cast(:user_id as uuid)
                  and s.type = 'website'
                  and j.status = 'running'
                order by j.updated_at desc nulls last
                limit 1
                """
            ),
            {"agent_id": str(agent_id), "user_id": str(user_id)},
        )
    ).mappings().first()
    running_crawl_bytes: int | None = None
    if running_crawl_row is not None and running_crawl_row["b"] is not None:
        running_crawl_bytes = int(running_crawl_row["b"])

    last_crawl_row = (
        await db.execute(
            text(
                """
                select (j.metrics->>'crawl_http_bytes')::bigint as b
                from public.indexing_jobs j
                join public.knowledge_sources s on s.id = j.knowledge_source_id
                where s.agent_id = cast(:agent_id as uuid)
                  and s.user_id = cast(:user_id as uuid)
                  and s.type = 'website'
                  and j.status = 'succeeded'
                  and j.metrics ? 'crawl_http_bytes'
                order by j.finished_at desc nulls last, j.updated_at desc
                limit 1
                """
            ),
            {"agent_id": str(agent_id), "user_id": str(user_id)},
        )
    ).mappings().first()
    last_crawl_bytes: int | None = None
    if last_crawl_row is not None and last_crawl_row["b"] is not None:
        last_crawl_bytes = int(last_crawl_row["b"])

    crawl_bytes_for_ui = running_crawl_bytes if running_crawl_bytes is not None else last_crawl_bytes

    await db.commit()

    return WebsiteUsageResponse(
        plan_slug=plan_slug,
        plan_name=plan_name,
        included_storage_bytes=included,
        used_storage_bytes=used,
        total_links=total_links,
        total_files=total_files,
        total_snippets=total_snippets,
        total_qa_pairs=total_qa_pairs,
        website_used_bytes=website_used,
        files_used_bytes=files_used,
        snippets_used_bytes=snippets_used,
        qa_used_bytes=qa_used,
        show_upgrade=show_upgrade,
        website_crawl_budget_bytes=remaining_total,
        website_crawl_last_job_bytes=crawl_bytes_for_ui,
    )
